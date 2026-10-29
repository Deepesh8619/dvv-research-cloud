"""Intake: load an authored questionnaire (.docx/PDF) into candidate questions.

This is the first, *deterministic* stage of questionnaire ingestion. It reads a
market-research questionnaire document and normalizes it into a flat list of
`DocumentBlock`s (paragraphs, headings, tables) in document order, then splits
that list into `QuestionChunk`s using purely rule-based / regex detection of the
question-label conventions MR authors use (Q5, S3, A1a., SECTION B, TERMINATE…).

No LLM calls happen here — a later stage feeds these chunks to an LLM to parse
them into structured survey logic. Grid questions are almost always authored as
Word tables (rows = options, columns = statements or vice-versa); we keep those
as a distinct `TableBlock`-shaped block so the later stage can recognize a grid
instead of seeing a flattened wall of text.

Design mirrors the tabular/spss adapters (architecture §adapters: every format
maps into a canonical model; downstream never sees the source format). Errors are
raised with clear messages, never swallowed.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

from src.intake import card_tables
from src.intake import taught_labels as taught
from src.intake.revisions import live_cell_text, live_text

logger = logging.getLogger(__name__)

BlockType = Literal["paragraph", "table", "heading"]


@dataclass
class DocumentBlock:
    """One structural unit of a questionnaire, format-independent.

    For table blocks, `text` holds a serialized representation and `rows` holds
    the real grid (list of row-lists of cell strings) so a grid question is
    never flattened into plain paragraph text.
    """

    block_type: BlockType
    text: str
    paragraph_index: int
    page_number: int | None = None  # PDF only; None for .docx
    is_bold: bool = False
    is_heading: bool = False
    rows: list[list[str]] | None = None  # populated only for block_type == "table"
    #: Text the document shows as taken out — struck through, or a Word
    #: tracked deletion. Not part of `text`, and kept so a reviewer can tell
    #: "the brand is missing" from "the brand was struck out in the document".
    removed: list[str] = field(default_factory=list)
    #: The document formatted this paragraph as a numbered or bulleted list
    #: item. Word keeps that numbering in the XML, so the number never appears
    #: in `text` — which is why an answer list can arrive as bare lines with no
    #: visible codes at all. It is the document's own statement that a line is
    #: an ITEM rather than a sentence, and on the three real client documents it
    #: separates the answer list from the scripter notes sitting beside it.
    is_list_item: bool = False
    #: The number Word draws in front of an auto-numbered paragraph ("2",
    #: "Q2.", "a)"), as the reader sees it; None when it has none.
    list_label: str | None = None
    #: The question this block opens, when the document's LAYOUT states it (a
    #: question card's label line). Trusted over any reading of the text,
    #: because "QGENDER." is a label only the layout can vouch for.
    label: str | None = None
    #: A grid the layout states outright, as ([(code, row)], [(code, column)]).
    grid: tuple[list[tuple[int, str]], list[tuple[int, str]]] | None = None


@dataclass
class QuestionChunk:
    """A candidate question: the blocks from one label line up to the next."""

    chunk_text: str
    candidate_qid: str | None
    source_blocks: list[DocumentBlock]
    provenance: str  # human-readable location, e.g. "page 3, para 12" or "para 47"


# --- Question-label detection (rule-based only) -----------------------------

# Q5, S3, A1a., Q12:, D2), P1., T3. — a short letter prefix + number, optional
# suffix letter, terminated by a separator so plain words like "Sales" don't
# match. The prefix is ANY 1-3 letters, not a whitelist: a real client screener
# used P1-P3 (parent path) and T1-T5 (teen path), which a Q|S|A|C|D whitelist
# silently glued onto the preceding question (pilot finding, 2026-07).
# The sub-question suffix is EITHER case: Decipher client docs write "A1a.",
# but federal instruments write "D4A.", "D8A.", "BK49A." A lowercase-only
# suffix class silently glued every uppercase sub-question onto the preceding
# question (SHED 2014: D4A/D4B/D4C/D8A all lost). The suffix is kept as
# written — it is the variable name the data map will carry.
# A sub-item suffix "_1" (Q7_1) is part of the name, and a label may be followed
# by two spaces or a tab instead of punctuation ("Q7_1  NUM 1-60 How many...").
_QID_RE = re.compile(r"^\s*([A-Za-z]{1,3})(\d+)([A-Za-z]?(?:_\d+)?)(?:[.):]|\s{2,}\S|\t)")
# SECTION B, SECTION 2 — a labelled section boundary with no question number.
_SECTION_RE = re.compile(r"^\s*SECTION\s+[A-Z0-9]", re.IGNORECASE)
# Strong instruction/boundary markers even without a question-number prefix.
_INSTRUCTION_RE = re.compile(r"^\s*(?:TERMINATE|SCREEN\s*OUT|SKIP\s*TO)\b", re.IGNORECASE)
# Table-per-question format (real Google/GBK programming QREs): each question is
# one Word table whose serialized text carries a `Q Name` row naming the qid.
# The separator is `|` on the raw pipe-join path and `:` on the clean Q-Name
# serialization path (see _serialize_qname_table), so accept either. The qid is a
# named token (RESP_AGE, SS_EMP_STATUS), not a Q5-style label.
_QNAME_RE = re.compile(r"(?im)^\s*Q ?Name\s*[|:]\s*([A-Za-z][A-Za-z0-9_]*)")

# Boilerplate placeholder cell values that mean "nothing was filled in here",
# NOT real content — stripped whole-cell (case-insensitively) from Q-Name
# question tables so neither the LLM nor a human reading source_text sees
# template noise. Exact strings observed in real GBK programming QREs.
_BOILERPLATE_CELLS = frozenset(
    s.lower()
    for s in (
        "If Custom, specify:",
        "If not All, specify:",
        "If Yes, details:",
        "Provide additional scripting details here",
        "None",
    )
)
# Programmer notes authored ABOVE their question that belong WITH the following
# question table, not the preceding one.
_NOTE_RE = re.compile(r"^\s*SCRIPTER\s*NOTE\b", re.IGNORECASE)

# Some questionnaires — routinely the healthcare/pharma ones — number nothing.
# Every question simply opens with the base it is asked on:
#
#     [BASE: ALL] Thinking about the patients with TRD you treat...
#     [BASE: THOSE WHO PRESCRIBE SPRAVATO] What drives your choice?
#
# In a NUMBERED document that same tag is a directive belonging to the question
# above it, so treating it as a boundary would shred the documents that work
# today. It is promoted to a boundary only for a document that carries no
# numbering at all (see `_is_unnumbered`), where it is the only thing marking
# where one question ends and the next begins.
# "[BASE SELECTS 1-4 IN PREVIOUS QUESTION]" writes the tag with no colon. Only
# the colon form was accepted, so that question, its answers and the next two
# headings were swallowed into the question above it and never asked.
_BASE_TAG_RE = re.compile(r"^\s*\[\s*BASE\b(?:\s*[:\-]|\s+)", re.IGNORECASE)

# "Thank you! We appreciate your input!" — the end screen, not one more answer.
_CLOSING = re.compile(r"^\s*(?:thank\s+you|thanks)\b.{0,120}$", re.I)

# A screen the respondent READS, not answers: the document talking to them —
# a welcome, a hand-over, a transition into the next block. Recognised by what
# it does (addresses the respondent, introduces what comes next), not by a label.
_SCREEN = re.compile(
    r"^\s*(?:thank\s+you|thanks|hi\b|hello\b|welcome|"
    r"thinking\s+(?:now\s+)?(?:about|of)|moving\s+on|now,?\s+(?:we|i|please|let'?s|thinking)|"
    r"you\s+will\s+(?:now\s+)?(?:be|see)|these\s+(?:final|next|following|last)|"
    r"the\s+(?:next|following|final|last)\s+(?:few\s+)?(?:questions|section|part)|"
    r"next,?\s+(?:we|you|i)|please\s+(?:now\s+)?(?:hand|read|take|think|imagine|consider|pass)|"
    r"in\s+this\s+(?:section|part)|before\s+we\s+(?:begin|start)|"
    r"we(?:'d|\s+would)\s+(?:now\s+)?like|imagine|for\s+the\s+(?:next|following|rest)|"
    r"after\s+seeing|to\s+start|to\s+begin|finally|lastly)\b", re.I)


def _is_screen(block: DocumentBlock) -> bool:
    """A respondent-facing sentence that asks nothing: an information screen."""
    text = (block.text or "").strip()
    return (block.block_type == "paragraph" and bool(_SCREEN.match(text))
            and len(text.split()) >= 2 and "?" not in text and not _asks_a_question(block)
            and not text.startswith(("[", "<", "(")) and not _CAPS_HEADING.match(text))

# A Word list number drawn as a plain number: "7", "7.", "7)", "Q7.".
_LIST_NUMBER = re.compile(r"^\s*Q?(\d{1,3})\s*[.)]?\s*$", re.I)

# A section heading in an unnumbered document: ALL CAPS, two or more words.
_CAPS_HEADING = re.compile(
    r"^\s*(?!(?:NONE|OTHER|ALL OF|NOT SURE|DON|DO NOT KNOW|PREFER|REFUSED|N/A)\b)"
    r"(?=[^a-z]*$)(?=.*[A-Z]{2})"
    r"[A-Z0-9][A-Z0-9&/,:'’()\-–]*(?:\s+[A-Z0-9&/,:'’()\-–]+){1,11}\s*$"
)

# Bracketed directives NEST — real ones do, routinely:
#
#     [ASK IF (K2_b=1 (yes) OR [K14a=2 (no) or refused]) AND (D3 = 1 or 2)]
#
# A `\[[^\]]*\]` pattern stops at the inner `]`, so a nested directive matched
# nothing, fell through to the answer parser and became an option. Scanning by
# bracket depth reads them the way a person does.
def bracket_bodies(text: str) -> list[str]:
    """The top-level bracketed directive bodies in a line, nesting included."""
    bodies: list[str] = []
    depth = 0
    start = 0
    for i, ch in enumerate(text):
        if ch == "[":
            if depth == 0:
                start = i + 1
            depth += 1
        elif ch == "]" and depth:
            depth -= 1
            if depth == 0:
                bodies.append(text[start:i])
    return bodies


def strip_brackets(text: str) -> str:
    """`text` with every bracketed group removed, nesting included."""
    out: list[str] = []
    depth = 0
    for ch in text:
        if ch == "[":
            depth += 1
        elif ch == "]" and depth:
            depth -= 1
        elif depth == 0:
            out.append(ch)
    return "".join(out)


def is_directive_only(text: str) -> bool:
    """True when a line is nothing but bracketed directives — no prose."""
    return bool(text.strip()) and not strip_brackets(text).strip()
# Routing directives name WHO gets asked. They are authored immediately above
# the question they govern, so a run of them trailing a chunk belongs to the
# NEXT question, not the one whose options just ended:
#
#     8. Retired
#     [SP] [IF D2 = 1]      <- D4A's routing and type, sitting in D2's chunk
#     D4A. Do you work full-time...
#
# Without this, D4A loses its base definition and D2 is handed a type marker
# that was never about D2 (Federal Reserve SHED 2014: 30+ questions affected).
# The trailing run only moves when it actually contains routing — a lone "[SP]"
# after a question's options is left where the author put it.
_ROUTING_DIRECTIVE_RE = re.compile(
    r"\[\s*(?:IF|ONLY\s+ASK|ASK\s+IF|ASK\s+ONLY)\b", re.IGNORECASE
)


# One identifier token; what follows it may be stray punctuation ("C_TEAM_S(ZE")
# but not more words, so a sentence never reads as an id.
_TOKEN_RE = re.compile(r"([A-Za-z][A-Za-z0-9_]*)(?=$|[^\w\s])")


def _rows_qid(rows: list[list[str]] | None) -> str | None:
    """The qid of a spec table, read from its rows: some row's label is one of
    the client's names for "question id" (Q Name, Variable, Question ID...,
    FIELD_ALIASES["qid"] in spec_table) and its value is one identifier token.
    Label wording lives in that one alias set, not here."""
    from src.engines.questionnaire.spec_table import FIELD_ALIASES, _norm_label

    for row in rows or ():
        cells = [c.strip() for c in row if c and c.strip()]
        if (
            len(cells) >= 2 and _norm_label(cells[0]) in FIELD_ALIASES["qid"]
            and (m := _TOKEN_RE.match(cells[1]))
        ):
            return m.group(1)
    return None


def _table_qid(block: DocumentBlock) -> str | None:
    """The token qid of a table-per-question block, or None for other blocks.

    A table is a self-contained question iff it has a `Q Name | <token>` row.
    Quota grids / outline tables have no such row and return None (they keep the
    default 'attach to preceding chunk' behavior)."""
    if block.block_type != "table":
        return None
    if qid := _rows_qid(block.rows):
        return qid
    m = _QNAME_RE.search(block.text)
    return m.group(1) if m else None


# A question named by an identifier on its own line (AGE, MAIN_BANK, LIFESTAGE
# (derived)), with no Q-number. Single ALL-CAPS token; only trusted when the
# document has several, so a lone "SCREENER" or "NONE" never opens a question.
_IDENT_LINE_RE = re.compile(r"^\s*([A-Z][A-Z0-9_]{1,39})(?:\s*\(.{0,60}\))?\s*$")
_IDENT_STOP = frozenset(
    "NONE OTHER TERMINATE CLOSE END SCREENER INTRO DEMOGRAPHICS YES NO ALL "
    "NA REFUSED THANKS CONTINUE ROWS COLUMNS".split()
)


def _ident_qid(block: DocumentBlock) -> str | None:
    if block.block_type != "paragraph":
        return None
    m = _IDENT_LINE_RE.match(block.text)
    return m.group(1) if m and m.group(1) not in _IDENT_STOP else None


def _has_ident_questions(blocks: list[DocumentBlock]) -> bool:
    """Several identifier lines AND no Q-number or spec table anywhere: in a
    numbered document an ALL-CAPS line is a heading or a scripter note."""
    return sum(1 for b in blocks if _ident_qid(b)) >= 3 and not any(
        _QID_RE.match(b.text) or _table_qid(b) is not None for b in blocks
    )


def _clean_cell(value: str) -> str:
    """A cell's stripped content, or '' if it is only a boilerplate placeholder."""
    v = value.strip()
    return "" if v.lower() in _BOILERPLATE_CELLS else v


def _row_value(cells: list[str]) -> str:
    """Collapse a row's value cells: drop merged-cell duplicates and boilerplate.

    Merged Word cells repeat the same value across every column; we keep only
    the distinct real values (in order). Genuinely distinct cells — a rare true
    multi-column row — are preserved, joined by ' | '."""
    seen: list[str] = []
    for cell in cells:
        c = _clean_cell(cell)
        if c and c not in seen:
            seen.append(c)
    return " | ".join(seen)


def _serialize_qname_table(rows: list[list[str]]) -> str:
    """Serialize a Q-Name question table as clean `Label: value` lines.

    Naive ' | '.join(row) turns a merged-cell row into
    `Q Name | RESP_AGE | RESP_AGE | RESP_AGE` — repetitive noise that confuses
    the LLM extractor and is unpleasant to review. Here each row becomes its
    label (col A) plus the de-duplicated, boilerplate-stripped value, so both the
    model and a human see clean content. Generic (non-Q-Name) tables keep the
    untouched pipe-join path — this only runs when a `Q Name` row is present.
    """
    lines: list[str] = []
    for row in rows:
        if not row:
            continue
        label = row[0].strip().rstrip(":").strip()
        value = _row_value(row[1:]) if len(row) > 1 else ""
        if not label:
            # No label cell (a 'Response list:' header, or an option value whose
            # code cell merged away): emit the value alone.
            if value:
                lines.append(value)
        elif value == label:
            # All cells identical (section header, question text): emit once.
            lines.append(label)
        elif value:
            # Labelled row with real content: `Label: value` (e.g. option
            # `99: Prefer not to answer`, `Q Name: RESP_AGE`).
            lines.append(f"{label}: {value}")
        else:
            # Labelled but empty (e.g. `Among who` on a question asked of all,
            # or a blank `Q type`): make the emptiness explicit for the reader.
            lines.append(f"{label}: (blank)")
    return "\n".join(lines)


def _is_unnumbered(blocks: list[DocumentBlock]) -> bool:
    """True when nothing in the document names a question.

    A document with even one `Q5.` label or one `Q Name:` table is numbered, and
    its base tags stay directives. Only a document with no names anywhere has to
    fall back to base tags to find its question boundaries.
    """
    return not _has_ident_questions(blocks) and not any(
        _QID_RE.match(b.text) or _table_qid(b) is not None for b in blocks
    )


def _detect_label(text: str, base_tags_open_questions: bool = False) -> tuple[bool, str | None]:
    """Return (is_boundary, candidate_qid) for a block's leading text.

    A question label yields its qid (e.g. "Q5"); a SECTION/instruction boundary
    is a boundary with no qid. Anything else is not a boundary.

    `base_tags_open_questions` is set only for an unnumbered document, where a
    leading `[BASE: ...]` tag is what opens each question.
    """
    m = _QID_RE.match(text)
    if m:
        # Normalize the letter prefix to uppercase (q5a -> Q5a, pn2 -> PN2);
        # a trailing sub-question letter stays lowercase as written.
        prefix, number, suffix = m.groups()
        return True, prefix.upper() + number + suffix
    if _SECTION_RE.match(text) or _INSTRUCTION_RE.match(text):
        return True, None
    if base_tags_open_questions and _BASE_TAG_RE.match(text):
        return True, None
    return False, None


def _provenance(block: DocumentBlock) -> str:
    if block.page_number is not None:
        return f"page {block.page_number}, para {block.paragraph_index}"
    return f"para {block.paragraph_index}"


# --- .docx loading ----------------------------------------------------------


def load_docx(path: str | Path) -> list[DocumentBlock]:
    """Load a .docx into ordered DocumentBlocks (paragraphs, headings, tables).

    Paragraphs and tables are emitted in true document order. Bold is detected
    from a paragraph's runs; a paragraph is a heading if it uses a Word "Heading"
    style. Tables are kept as `rows` (list of row-lists) and never flattened.
    """
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency guaranteed by requirements
        raise RuntimeError(
            "python-docx is required to load .docx questionnaires; "
            "install it (pip install python-docx>=1.0)"
        ) from exc

    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"questionnaire .docx not found: {path}")

    document = Document(str(path))
    body = document.element.body
    blocks: list[DocumentBlock] = []
    para_index = 0
    struck_out: list[str] = []
    # Each table's distinct cells as lines, by its position in `blocks`, so a
    # document laid out as question cards can be re-read once its layout is known.
    table_cells: dict[int, list[list[list[str]]]] = {}

    # Map child XML elements back to python-docx objects to walk in order.
    para_map = {p._p: p for p in document.paragraphs}
    table_map = {t._tbl: t for t in document.tables}

    from src.intake.word_numbering import ListNumbering

    numbering = ListNumbering(document)
    for child in body.iterchildren():
        if child in para_map:
            para: Paragraph = para_map[child]
            if _is_contents_entry(para):
                # "KEY METRICS ....... 22" is the document's table of contents,
                # not a question. Read as one, it became question SCREENER8
                # "KEY METRICS22" (heading + page number) on a live GBK file.
                continue
            text, removed = live_text(para)
            # Advanced for every paragraph Word still numbers — an empty one
            # included — and not for one deleted outright.
            list_label = numbering.label(para) if (text or not removed) else None
            if not text:
                # A line that is ENTIRELY struck out is gone from the survey —
                # in the Vivo wave-2 document that is a whole question, which
                # was being extracted and would have been fielded.
                if removed:
                    struck_out.extend(removed)
                continue
            is_bold = _para_is_bold(para)
            is_heading = (para.style.name or "").lower().startswith("heading")
            is_list_item = _is_list_paragraph(para)
            blocks.append(
                DocumentBlock(
                    block_type="heading" if is_heading else "paragraph",
                    text=text,
                    paragraph_index=para_index,
                    page_number=None,
                    is_bold=is_bold,
                    is_heading=is_heading,
                    removed=removed,
                    is_list_item=is_list_item,
                    list_label=list_label,
                )
            )
            para_index += 1
        elif child in table_map:
            table: Table = table_map[child]
            cells = [[live_cell_text(cell) for cell in row.cells] for row in table.rows]
            rows = [[text for text, _ in row] for row in cells]
            table_removed = [gone for row in cells for _, removes in row for gone in removes]
            # A Q-Name question table gets clean `Label: value` serialization
            # (merged-cell duplication + boilerplate removed); generic tables
            # (quota grids, outlines) keep the untouched pipe-join.
            clean = _serialize_qname_table(rows)
            serialized = clean if _rows_qid(rows) else "\n".join(
                " | ".join(r) for r in rows
            )
            table_cells[len(blocks)] = [
                [
                    [x for x in live_cell_text(cell)[0].split("\n") if x.strip()]
                    for cell in _distinct(row.cells)
                ]
                for row in table.rows
            ]
            blocks.append(
                DocumentBlock(
                    block_type="table",
                    text=serialized,
                    paragraph_index=para_index,
                    page_number=None,
                    is_bold=False,
                    is_heading=False,
                    rows=rows,
                    removed=table_removed,
                )
            )
            para_index += 1

    blocks = _read_cards(blocks, table_cells)
    if struck_out:
        logger.info(
            "%s: %d struck-out fragments left out of the questionnaire",
            path.name, len(struck_out),
        )
    return blocks


def _distinct(cells) -> list:
    """A row's cells with a merged cell once (Word repeats it per column)."""
    seen: list = []
    for cell in cells:
        if all(cell._tc is not other._tc for other in seen):
            seen.append(cell)
    return seen


def _read_cards(
    blocks: list[DocumentBlock], table_cells: dict[int, list[list[list[str]]]]
) -> list[DocumentBlock]:
    """Rewrite a question-card document's tables as the lines they stand for.

    Decided for the WHOLE document from its layout: only when several tables
    are question cards (card_tables.MIN_CARDS) is any table rewritten, and then
    every card plus the single-column heading/intro tables around them.
    Everything else in the document, and every other document, is untouched.
    """
    cards = {i for i, cells in table_cells.items() if card_tables.label_of(cells)}
    if len(cards) < card_tables.MIN_CARDS:
        return blocks
    out: list[DocumentBlock] = []
    for i, block in enumerate(blocks):
        cells = table_cells.get(i)
        lines, grid = (
            card_tables.card_lines(cells) if i in cards
            else (card_tables.plain_lines(cells), None) if cells else ([], None)
        )
        if not lines:
            out.append(block)
            continue
        # Every line keeps the table's own position, so "para 202" still
        # opens the card in the document view.
        out += [
            DocumentBlock(block_type="paragraph", text=line,
                          paragraph_index=block.paragraph_index, removed=list(block.removed),
                          label=card_tables.label_of(cells) if i in cards else None,
                          grid=grid)
            if n == 0 else
            DocumentBlock(block_type="paragraph", text=line,
                          paragraph_index=block.paragraph_index)
            for n, line in enumerate(lines)
        ]
    return out


_CONTENTS_TITLE = re.compile(r"^\s*(table of contents|contents)\s*$", re.I)


def _is_contents_entry(paragraph) -> bool:
    """A table-of-contents line: Word's TOC styles ("TOC 1", "toc 2"), a link to
    a `_Toc` heading bookmark, a PAGEREF/TOC field, or the "Contents" title.
    Any one is enough — each is how Word writes its own contents, never how a
    questionnaire writes a question."""
    style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
    if style.startswith("toc") or style in ("table of contents", "contents"):
        return True
    p = paragraph._p
    if p.xpath('.//w:hyperlink[starts-with(@w:anchor, "_Toc")]'):
        return True
    if any(("PAGEREF _Toc" in t or t.strip().startswith("TOC "))
           for t in p.xpath(".//w:instrText/text()")):
        return True
    return bool(_CONTENTS_TITLE.match(paragraph.text or "")) and len(paragraph.text) < 30


def _is_list_paragraph(paragraph) -> bool:
    """Did the document format this paragraph as a list item?

    Two ways Word records it, and a document uses one or the other. The three
    real client questionnaires carry `w:numPr` directly on the paragraph; a
    document built from Word's own "List Number" / "List Bullet" styles carries
    it on the STYLE instead, and checking only the paragraph found nothing at
    all there.
    """
    if paragraph._p.xpath(".//w:numPr"):
        return True
    style = paragraph.style
    while style is not None:
        name = (style.name or "").lower()
        if name.startswith(("list number", "list bullet", "list paragraph")):
            return True
        if getattr(style, "element", None) is not None and style.element.xpath(
            ".//w:numPr"
        ):
            return True
        style = getattr(style, "base_style", None)
    return False


def _para_is_bold(paragraph) -> bool:
    """A paragraph is 'bold' if it has runs and every non-empty run is bold."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    if not runs:
        return False
    return all(bool(r.bold) for r in runs)


# --- PDF loading ------------------------------------------------------------

# Blocks whose left edges cluster into >1 group wider than this fraction of the
# page width are treated as a multi-column layout.
_COLUMN_GAP_FRACTION = 0.15


def load_pdf(path: str | Path) -> list[DocumentBlock]:
    """Load a PDF into ordered DocumentBlocks with reading-order awareness.

    Naive top-to-bottom extraction scrambles multi-column layouts, so per page
    we detect columns by clustering text blocks on their left x-coordinate and
    sort column-then-row. A block is flagged as a heading when its dominant font
    size is clearly larger than the page's body size (or it is bold).
    """
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - dependency guaranteed by requirements
        raise RuntimeError(
            "PyMuPDF (fitz) is required to load PDF questionnaires; "
            "install it (pip install PyMuPDF>=1.24)"
        ) from exc

    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"questionnaire PDF not found: {path}")

    blocks: list[DocumentBlock] = []
    para_index = 0
    with fitz.open(str(path)) as doc:
        for page_number, page in enumerate(doc, start=1):
            page_blocks = _extract_page_blocks(page)
            body_size = _body_font_size(page_blocks)
            for pb in page_blocks:
                # One DocumentBlock per LINE, not per PDF block. A PDF "block"
                # is a visual paragraph group, so a single one routinely spans
                # several logical lines — the previous question's answer list
                # followed by the next question's label:
                #
                #   "1. Yes  0. No  [SP] S2. Which best describes your housing…"
                #
                # `_detect_label` only matches at position 0, so every label
                # buried mid-block was invisible and the whole document
                # collapsed into a handful of chunks (measured on the Federal
                # Reserve SHED 2014 questionnaire: 268 blocks -> 4 chunks for a
                # ~100-question survey). Lines exist in the PDF; joining them
                # away was throwing the boundaries out. Ordering still happens
                # at block level, so column detection is unchanged.
                for line in pb["lines"]:
                    is_heading = line["is_bold"] or (
                        body_size is not None and line["size"] > body_size * 1.15
                    )
                    blocks.append(
                        DocumentBlock(
                            block_type="heading" if is_heading else "paragraph",
                            text=line["text"],
                            paragraph_index=para_index,
                            page_number=page_number,
                            is_bold=line["is_bold"],
                            is_heading=is_heading,
                        )
                    )
                    para_index += 1
    return blocks


def _extract_page_blocks(page) -> list[dict]:
    """Return this page's text blocks sorted in column-then-row reading order.

    Each dict carries: text, x0, y0, size (dominant font size), is_bold.
    """
    raw = page.get_text("dict")
    page_width = page.rect.width or 1.0
    parsed: list[dict] = []
    for block in raw.get("blocks", []):
        if block.get("type", 0) != 0:  # skip image blocks
            continue
        line_items: list[dict] = []
        sizes: list[float] = []
        bold_flags: list[bool] = []
        for line in block.get("lines", []):
            line_parts: list[str] = []
            line_sizes: list[float] = []
            line_bold: list[bool] = []
            for span in line.get("spans", []):
                span_text = span.get("text", "")
                if not span_text.strip():
                    continue
                line_parts.append(span_text)
                line_sizes.append(float(span.get("size", 0.0)))
                flags = int(span.get("flags", 0))
                # PyMuPDF font flag bit 4 (value 16) marks bold.
                line_bold.append(bool(flags & 16))
            line_text = " ".join(line_parts).strip()
            if not line_text:
                continue
            line_items.append(
                {
                    "text": line_text,
                    "size": max(line_sizes) if line_sizes else 0.0,
                    "is_bold": any(line_bold),
                }
            )
            sizes.extend(line_sizes)
            bold_flags.extend(line_bold)
        if not line_items:
            continue
        x0, y0 = block["bbox"][0], block["bbox"][1]
        parsed.append(
            {
                "text": " ".join(li["text"] for li in line_items),
                "lines": line_items,
                "x0": x0,
                "y0": y0,
                "size": max(sizes) if sizes else 0.0,
                "is_bold": any(bold_flags),
            }
        )

    return _sort_reading_order(parsed, page_width)


def _sort_reading_order(parsed: list[dict], page_width: float) -> list[dict]:
    """Sort blocks column-then-row when a multi-column layout is detected."""
    if len(parsed) < 2:
        return parsed

    columns = _cluster_columns([b["x0"] for b in parsed], page_width)
    if len(columns) < 2:
        # Single column: plain top-to-bottom (then left-to-right) order.
        return sorted(parsed, key=lambda b: (round(b["y0"]), b["x0"]))

    def column_index(x0: float) -> int:
        return min(range(len(columns)), key=lambda i: abs(x0 - columns[i]))

    # Column first (left-to-right), then row (top-to-bottom) within a column.
    return sorted(parsed, key=lambda b: (column_index(b["x0"]), round(b["y0"])))


def _cluster_columns(x_positions: list[float], page_width: float) -> list[float]:
    """Cluster left x-coordinates into column anchors (their mean x).

    Two anchors closer than `_COLUMN_GAP_FRACTION` of the page width are merged,
    so incidental indentation doesn't read as a second column.
    """
    gap = page_width * _COLUMN_GAP_FRACTION
    anchors: list[list[float]] = []
    for x in sorted(x_positions):
        placed = False
        for group in anchors:
            if abs(x - (sum(group) / len(group))) <= gap:
                group.append(x)
                placed = True
                break
        if not placed:
            anchors.append([x])
    return [sum(group) / len(group) for group in anchors]


def _body_font_size(page_blocks: list[dict]) -> float | None:
    """Most common (rounded) font size on the page — its body-text size."""
    sizes = [round(b["size"]) for b in page_blocks if b["size"] > 0]
    if not sizes:
        return None
    return float(max(set(sizes), key=sizes.count))


# --- Segmentation -----------------------------------------------------------


def _trailing_routing_directives(current: list[DocumentBlock]) -> list[DocumentBlock]:
    """Pop a trailing run of directive-only blocks that carries routing.

    Mutates `current`, returning the popped blocks in original order so the
    caller can prepend them to the next chunk. Returns [] (and leaves `current`
    untouched) when the trailing run has no routing directive in it, or when the
    directives are all the chunk has — a chunk that is nothing but directives is
    a standalone programmer note and stays where it is.
    """
    tail: list[DocumentBlock] = []
    i = len(current)
    while i > 0 and is_directive_only(current[i - 1].text):
        i -= 1
    if i == 0 or i == len(current):
        return []
    tail = current[i:]
    if not any(_ROUTING_DIRECTIVE_RE.search(b.text) for b in tail):
        return []
    del current[i:]
    return tail


_LETTERED = re.compile(r"^\s*[A-H][.)]\s+\S")


def _asks_a_question(block: DocumentBlock) -> bool:
    """A paragraph worded as a question to a respondent, not a scripter note."""
    text = block.text.strip()
    return (
        block.block_type == "paragraph"
        and len(text) >= 40
        and "?" in text
        and not text.startswith("[")
    )


def segment_into_questions(
    blocks: list[DocumentBlock], taught_labels: list[str] | None = None
) -> list[QuestionChunk]:
    """Split ordered blocks into candidate questions on detected label lines.

    `taught_labels` are example labels owners taught ("QGENDER.", "Question 5:");
    each is applied only where this document writes several labels that way
    (taught_labels.py), so a label taught on one client reads the next.

    A chunk runs from one detected question-label / section / instruction line up
    to (not including) the next detected label, accumulating every block in
    between (including tables, which usually form that question's answer grid).
    Leading content before the first label is emitted as its own chunk with
    `candidate_qid=None` so nothing is silently dropped.
    """
    if taught_labels:
        blocks, _marked = taught.apply(blocks, taught_labels)
    chunks: list[QuestionChunk] = []
    current: list[DocumentBlock] = []
    current_qid: str | None = None
    seen_label = False
    # An unnumbered, base-tagged document. Its questions get sequential Q1..Qn
    # here rather than all arriving as UNKNOWN: a qid is the address the review
    # UI, approval and the data column all use, and "UNKNOWN_7" is not one a
    # reviewer can match back to the document they are reading.
    base_tags_open = _is_unnumbered(blocks)
    ident_named = _has_ident_questions(blocks)
    base_counter = 0
    # Word list numbers, when the questions carry them: the document's own
    # names. `listed` is the last one seen; `extra` counts unlisted questions
    # after it (named Q5_2, Q5_3 rather than stealing Q6 from the next one).
    listed: int | None = None
    extra = 0

    def next_qid(block: DocumentBlock) -> str:
        nonlocal base_counter, listed, extra
        number = _LIST_NUMBER.match(block.list_label or "")
        if number:
            listed, extra = int(number.group(1)), 0
            return f"Q{listed}"
        lettered = _LETTERED.match(block.text or "")
        if listed is not None and lettered:
            # "A. Thinking about…" under item 7 is 7a, as a reader cites it.
            return f"Q{listed}{block.text.strip()[0].lower()}"
        if listed is not None:
            extra += 1
            return f"Q{listed}_{extra + 1}"
        base_counter += 1
        return f"Q{base_counter}"

    def flush() -> None:
        if not current:
            return
        if (
            base_tags_open and len(current) == 1 and current_qid is None
            and _CAPS_HEADING.match(current[0].text or "")
        ):
            # A section heading with nothing under it ("ADOPTION INTENT"). It
            # names a section; it is not a screen, and documents say outright
            # "do not show section headers".
            return
        # An identifier line names the question (candidate_qid); the wording
        # is what follows, so the id is not also read as the question text.
        shown = list(current)
        if ident_named and current_qid:
            for i, b in enumerate(shown):
                if _ident_qid(b) == current_qid:
                    del shown[i]
                    break
        chunks.append(
            QuestionChunk(
                chunk_text="\n".join(b.text for b in shown).strip(),
                candidate_qid=current_qid,
                source_blocks=list(current),
                provenance=_provenance(current[0]),
            )
        )

    screens = 0

    def starts_question(block: DocumentBlock) -> bool:
        return block.label is not None or _table_qid(block) is not None \
            or _detect_label(block.text, base_tags_open)[0] \
            or (ident_named and _ident_qid(block) is not None)

    def before_a_question(i: int) -> bool:
        """The next block that is not a scripter note or another screen starts a question."""
        for later in blocks[i + 1:i + 6]:
            text = (later.text or "").strip()
            if starts_question(later):
                return True
            if _is_screen(later) or text.startswith(("[", "<")) or _NOTE_RE.match(text) \
                    or _CAPS_HEADING.match(text):
                continue
            return False
        return False

    for i, block in enumerate(blocks):
        if _is_screen(block) and before_a_question(i) and not starts_question(block) \
                and (current_qid is not None or seen_label or current):
            # "Moving on…", "Thinking now about one brand, Bubble…", "Please now hand
            # your device to your teen": its own screen, read before the next question.
            # It used to ride at the end of the PREVIOUS question's text.
            if current_qid is not None and current and _is_screen(current[-1]) \
                    and current_qid.startswith(("INFO_", "INTRO")):
                current.append(block)  # a second sentence of the same screen
                continue
            flush()
            screens += 1
            current = [block]
            current_qid = "INTRO" if not seen_label and screens == 1 else f"INFO_{screens}"
            continue
        table_qid = _table_qid(block)
        if table_qid is not None:
            # A table-per-question block is its own chunk. Peel any immediately-
            # preceding SCRIPTER NOTE paragraphs off the current chunk so they
            # ride WITH this question rather than the previous one.
            notes: list[DocumentBlock] = []
            while (
                current
                and current[-1].block_type != "table"
                and _NOTE_RE.match(current[-1].text)
            ):
                notes.insert(0, current.pop())
            flush()
            current = notes + [block]
            current_qid = table_qid
            seen_label = True
            continue

        is_boundary, qid = (
            (True, block.label) if block.label else _detect_label(block.text, base_tags_open)
        )
        if not is_boundary and ident_named and (ident := _ident_qid(block)):
            is_boundary, qid = True, ident
        if (
            not is_boundary and base_tags_open and current
            and current[-1].block_type == "table"
            and (_CAPS_HEADING.match(block.text or "") or _CLOSING.match(block.text or ""))
        ):
            # "PRODUCT X PROFILE" after an answer table: a section heading, not
            # one more answer. Read as answers, it and the next heading were
            # offered to respondents as choices.
            is_boundary = True
        if is_boundary and qid is None and base_tags_open and _BASE_TAG_RE.match(block.text):
            qid = next_qid(block)
        elif (
            not is_boundary and base_tags_open and current and _asks_a_question(block)
            and (current[-1].block_type == "table" or _LETTERED.match(block.text))
        ):
            # Also "A. Thinking about your TRD patients, what percentage...?":
            # lettered sub-questions under an intro ("For the following
            # questions...") were read as the previous question's answers.
            # A question whose [BASE:] tag the author forgot. A question's
            # answer table ends that question, so a question sentence straight
            # after one starts the next — a scripter reads it that way. Without
            # this, the TRD referrer study's settings question rode inside the
            # previous single-choice question and was never asked.
            is_boundary, qid = True, next_qid(block)
        elif (
            not is_boundary and base_tags_open and listed is not None
            and (number := _LIST_NUMBER.match(block.list_label or ""))
            and int(number.group(1)) == listed + 1
        ):
            # The next number of the question list ("15  Please describe why
            # you…", "7  For the following questions…"): a question, even with
            # no base tag and no question mark. Only the NEXT number: an answer
            # list numbered "1. Yes / 2. No" never continues the question list.
            is_boundary, qid = True, next_qid(block)
        if is_boundary:
            carried = _trailing_routing_directives(current)
            flush()
            current = carried + [block]
            current_qid = qid
            seen_label = True
        else:
            if not seen_label and not current:
                # Preamble before any label: start a headless (qid=None) chunk.
                current = [block]
                current_qid = None
            else:
                current.append(block)

    flush()
    return chunks
