"""The questionnaire as the client wrote it, rendered so a reviewer can look.

The parser reads a document for its MEANING and throws the rest away: a brand
struck out in red becomes an entry in `removed`, a highlighted note becomes
plain text, and a comment written in blue beside an option becomes the same
characters as everything else. That is right for extraction and wrong for
review, because the colour was carrying information:

    1: Google Ads
    [IF GOOGLE ADS CURRENT USER]        <- highlighted, in the client's yellow

A highlight after an option label is how these documents say "there is a
hidden variable behind this one". A reviewer who can see the highlight
connects it in a second. A reviewer reading the parser's flattened text
cannot see that anything was ever there.

So this renders the SAME walk the loader makes, keeping what the loader
discards: colour, highlight, bold, italic, underline, strikethrough, and
Word's tracked insertions and deletions. It is a reading view, never an
input — nothing here feeds extraction.

**Anchors.** Each block carries the id `b{n}` where n is the loader's
`paragraph_index`, so a question whose provenance says "para 47" scrolls to
`#b47`. The subtlety worth stating: the loader SKIPS a paragraph that is
entirely struck out, and skips it before incrementing the counter — but this
view must still show it, because "the whole question was struck out" is
exactly what a reviewer is looking for. So a struck-out-only paragraph is
rendered WITHOUT an id and does not advance the count. Anchors stay aligned
with the extracted questions; the deleted line is still on screen.

**Untrusted input.** The document arrives from a client. Every character of
its text is escaped, and the only attributes emitted are ones built here from
values this module validates — a colour is used only if it is six hex digits.
No markup from the document is ever passed through.
"""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path

from src.intake.revisions import W, _is_deleted, _is_struck

#: Word writes colours as RRGGBB, and "auto" for the default.
_HEX = re.compile(r"^[0-9A-Fa-f]{6}$")
#: A highlight is a named colour from Word's own fixed list.
_HIGHLIGHTS = {
    "yellow", "green", "cyan", "magenta", "blue", "red", "darkBlue", "darkCyan",
    "darkGreen", "darkMagenta", "darkRed", "darkYellow", "darkGray", "lightGray", "black",
}
#: Near-white text on white is invisible; the client meant it to be read.
_TOO_PALE = {"ffffff", "fefefe"}


@dataclass
class RenderedDocument:
    """The document as HTML, plus how many anchored blocks it holds."""

    html: str
    n_blocks: int

    def to_dict(self) -> dict:
        return {"html": self.html, "n_blocks": self.n_blocks}


def _on(properties, tag: str) -> bool:
    """A boolean run property, respecting an explicit w:val="false"."""
    if properties is None:
        return False
    element = properties.find(f"{W}{tag}")
    return element is not None and element.get(f"{W}val") not in ("false", "0")


def _run_html(run) -> str:
    """One run, with the formatting the parser drops.

    Deleted text lives in `w:delText` rather than `w:t`, so it contributes
    nothing to a normal reading and has to be asked for by name.
    """
    text = "".join(node.text or "" for node in run.findall(f"{W}t"))
    gone = "".join(node.text or "" for node in run.findall(f"{W}delText"))
    body = html.escape(text or gone)
    if not body:
        return ""

    properties = run.find(f"{W}rPr")
    classes = []
    styles = []

    if gone or _is_deleted(run):
        classes.append("dvv-del")
    elif _is_struck(run):
        classes.append("dvv-struck")
    if _in_insertion(run):
        classes.append("dvv-ins")
    if _on(properties, "b"):
        classes.append("dvv-b")
    if _on(properties, "i"):
        classes.append("dvv-i")
    if _on(properties, "u") or (
        properties is not None and properties.find(f"{W}u") is not None
    ):
        classes.append("dvv-u")

    if properties is not None:
        colour = properties.find(f"{W}color")
        value = (colour.get(f"{W}val") or "") if colour is not None else ""
        if _HEX.match(value) and value.lower() not in _TOO_PALE:
            styles.append(f"color:#{value.lower()}")
        highlight = properties.find(f"{W}highlight")
        name = (highlight.get(f"{W}val") or "") if highlight is not None else ""
        if name in _HIGHLIGHTS:
            classes.append(f"dvv-hl-{name}")

    if not classes and not styles:
        return body
    attrs = ""
    if classes:
        attrs += f' class="{" ".join(classes)}"'
    if styles:
        attrs += f' style="{";".join(styles)}"'
    return f"<span{attrs}>{body}</span>"


def _in_insertion(run) -> bool:
    """A run inside a Word tracked insertion."""
    parent = run.getparent()
    while parent is not None:
        if parent.tag == f"{W}ins":
            return True
        if parent.tag in (f"{W}p", f"{W}tc"):
            return False
        parent = parent.getparent()
    return False


def _paragraph_html(paragraph) -> tuple[str, bool]:
    """(inner HTML, whether anything survives as live text)."""
    parts = [_run_html(run) for run in paragraph._p.iter(f"{W}r")]
    live = any(
        "".join(n.text or "" for n in run.findall(f"{W}t")).strip()
        and not (_is_struck(run) or _is_deleted(run))
        for run in paragraph._p.iter(f"{W}r")
    )
    return "".join(parts), live


def _row_html(row) -> str:
    """One table row, with merged cells merged rather than repeated.

    python-docx hands back the SAME cell object once per grid column it spans,
    so a two-column row of a merged header arrived as five identical <td>s and
    the pane showed every value two and three times over. The document's real
    shape is the thing a reviewer is here to see, so consecutive cells backed
    by the same XML element become one <td colspan=n>.
    """
    cells: list[str] = []
    span = 0
    previous = None
    for cell in row.cells:
        if previous is not None and cell._tc is previous._tc:
            span += 1
            continue
        if previous is not None:
            cells.append(_cell_html(previous, span))
        previous, span = cell, 1
    if previous is not None:
        cells.append(_cell_html(previous, span))
    return "".join(cells)


def _cell_html(cell, span: int) -> str:
    inner = "".join(_paragraph_html(p)[0] for p in cell.paragraphs)
    width = f' colspan="{span}"' if span > 1 else ""
    return f"<td{width}>{inner or '&nbsp;'}</td>"


def render_docx(path: str | Path) -> RenderedDocument:
    """Render a .docx for reading, anchored to the loader's block indices."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    body = document.element.body
    para_map = {p._p: p for p in document.paragraphs}
    table_map = {t._tbl: t for t in document.tables}

    out: list[str] = []
    index = 0
    for child in body.iterchildren():
        if child in para_map:
            paragraph: Paragraph = para_map[child]
            inner, live = _paragraph_html(paragraph)
            if not inner.strip():
                continue
            heading = (paragraph.style.name or "").lower().startswith("heading")
            tag = "h3" if heading else "p"
            if live:
                # Anchored: the loader emitted a block here too.
                out.append(f'<{tag} id="b{index}" class="dvv-blk">{inner}</{tag}>')
                index += 1
            else:
                # Struck out entirely — the loader dropped it and did NOT count
                # it. Shown, because that is what a reviewer is looking for,
                # and left unanchored so every later anchor still lines up.
                out.append(f'<{tag} class="dvv-blk dvv-gone">{inner}</{tag}>')
        elif child in table_map:
            table: Table = table_map[child]
            rows = [f"<tr>{_row_html(row)}</tr>" for row in table.rows]
            out.append(
                f'<table id="b{index}" class="dvv-blk dvv-tbl">{"".join(rows)}</table>'
            )
            index += 1

    return RenderedDocument(html="\n".join(out), n_blocks=index)
