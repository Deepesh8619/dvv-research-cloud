"""Excel reports built from instructions: finished, or refreshable by formula.

Finished: the tables as the engine computed them — bases, column %, significance
letters, nets in bold, optional counts and native Excel charts — one sheet per
table or all on one sheet, behind a contents page.

Refreshable: the same layout, but every percentage and base is a COUNTIFS
formula over a Data sheet. The client pastes next wave's respondents into Data
and the report recalculates in any Excel, with no macro and nothing to enable.
Significance letters and means need the engine, so they are in the finished
version only; the refreshable one says so on its contents page.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from itertools import product
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INK = "16202C"
MUTED = "5B6878"
HEAD_FILL = PatternFill("solid", fgColor="EEF2F6")
NET_FILL = PatternFill("solid", fgColor="F6F8FA")
TITLE = Font(bold=True, size=13, color=INK)
BOLD = Font(bold=True, color=INK)
SMALL = Font(size=9, color=MUTED)
LETTER = Font(size=9, color="1F6F78", bold=True)
#: Rows a refreshable Data sheet's formulas read: room for later waves.
DATA_ROWS = 100_000


@dataclass
class ReportStyle:
    title: str
    layout: str = "sheet_per_table"  # or "one_sheet"
    decimals: int = 0
    show_counts: bool = False
    significance: bool = True
    bold_nets: bool = True
    charts: bool = False
    notes: list[str] = field(default_factory=list)


def _pct_format(decimals: int) -> str:
    return "0%" if decimals <= 0 else "0." + "0" * decimals + "%"


def _sheet_name(i: int, stub: str, taken: set[str]) -> str:
    name = f"T{i} {stub}"[:31].replace("/", "_").replace(":", "_")
    while name in taken:
        name = name[:28] + f"_{len(taken)}"
    taken.add(name)
    return name


def _contents(wb: Workbook, style: ReportStyle, meta: dict, entries: list[tuple[str, str, str]]):
    ws = wb.active
    ws.title = "Contents"
    ws["A1"] = style.title
    ws["A1"].font = Font(bold=True, size=15, color=INK)
    ws["A2"] = f"{meta.get('project', '')} · generated {datetime.now():%d %b %Y %H:%M}"
    ws["A2"].font = SMALL
    row = 4
    for note in [*meta.get("notes", []), *style.notes]:
        ws.cell(row=row, column=1, value=note).font = SMALL
        row += 1
    row += 1
    for col, head in enumerate(("Table", "Question", "Sheet"), start=1):
        ws.cell(row=row, column=col, value=head).font = BOLD
    for table_id, label, sheet in entries:
        row += 1
        ws.cell(row=row, column=1, value=table_id)
        ws.cell(row=row, column=2, value=label)
        link = ws.cell(row=row, column=3, value=sheet)
        link.hyperlink = f"#'{sheet}'!A1"
        link.font = Font(color="1F6F78", underline="single")
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 22


def _chart(ws, title: str, first_row: int, last_row: int, n_cols: int, anchor: str):
    chart = BarChart()
    chart.type = "bar"
    chart.title = title[:80]
    chart.y_axis.numFmt = "0%"
    chart.y_axis.scaling.min = 0  # never a truncated axis
    data = Reference(ws, min_col=2, max_col=1 + n_cols, min_row=first_row - 1, max_row=last_row)
    cats = Reference(ws, min_col=1, min_row=first_row, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 7.5, 15
    ws.add_chart(chart, anchor)


# --- finished -------------------------------------------------------------------


def _write_static_table(ws, top: int, table_id: str, r, style: ReportStyle) -> int:
    """One computed table at row `top`. Returns the next free row."""
    fmt = _pct_format(style.decimals)
    n_cols = len(r.columns)
    ws.cell(row=top, column=1, value=f"{table_id} · {r.stub_label or r.stub}").font = TITLE
    if r.filter_desc:
        ws.cell(row=top + 1, column=1, value=f"Filter: {r.filter_desc}").font = SMALL
    head = top + 2
    ws.cell(row=head, column=1, value=r.stub).font = BOLD
    for j, col in enumerate(r.columns, start=2):
        heading = col.label + (f" ({col.letter})" if col.letter else "")
        c = ws.cell(row=head, column=j, value=heading)
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center")
    base = head + 1
    ws.cell(row=base, column=1, value="Base (n)").font = SMALL
    for j, col in enumerate(r.columns, start=2):
        ws.cell(row=base, column=j, value=round(col.base)).font = SMALL
    row = base + 1
    first_pct = row
    sig_on = style.significance and r.sig
    for i, label in enumerate(r.row_labels):
        ws.cell(row=row, column=1, value=label)
        for j in range(n_cols):
            c = ws.cell(row=row, column=j + 2, value=float(r.col_pct[i][j]) / 100)
            c.number_format = fmt
        row += 1
        if sig_on and any(r.sig[i]):
            for j in range(n_cols):
                if r.sig[i][j]:
                    ws.cell(row=row, column=j + 2, value=r.sig[i][j]).font = LETTER
            row += 1
    last_pct = row - 1
    if r.net_pct is not None:
        for i, label in enumerate(r.net_labels):
            ws.cell(row=row, column=1, value=label).font = BOLD if style.bold_nets else None
            for j in range(n_cols):
                c = ws.cell(row=row, column=j + 2, value=float(r.net_pct[i][j]) / 100)
                c.number_format = fmt
                if style.bold_nets:
                    c.font, c.fill = BOLD, NET_FILL
            row += 1
    if r.means is not None:
        ws.cell(row=row, column=1, value="Mean")
        for j in range(n_cols):
            ws.cell(row=row, column=j + 2, value=round(float(r.means[j]), 2))
        row += 1
    if style.show_counts:
        ws.cell(row=row, column=1, value="Counts").font = SMALL
        row += 1
        for i, label in enumerate(r.row_labels):
            ws.cell(row=row, column=1, value=label).font = SMALL
            for j in range(n_cols):
                ws.cell(row=row, column=j + 2, value=round(float(r.counts[i][j]))).font = SMALL
            row += 1
    if style.charts and not sig_on:
        _chart(ws, r.stub_label or r.stub, first_pct, last_pct, n_cols,
               f"{get_column_letter(n_cols + 3)}{top}")
    elif style.charts:  # letter rows sit between the % rows: chart a clean copy
        _chart_copy(ws, r, top, n_cols, fmt)
    return row + 2


def _chart_copy(ws, r, top: int, n_cols: int, fmt: str):
    """Charted from a compact copy of the % rows, beside the table (letters excluded)."""
    left = n_cols + 3
    ws.cell(row=top + 2, column=left, value=r.stub).font = SMALL
    for j, col in enumerate(r.columns, start=1):
        ws.cell(row=top + 2, column=left + j, value=col.label).font = SMALL
    for i, label in enumerate(r.row_labels):
        ws.cell(row=top + 3 + i, column=left, value=label).font = SMALL
        for j in range(n_cols):
            c = ws.cell(row=top + 3 + i, column=left + 1 + j, value=float(r.col_pct[i][j]) / 100)
            c.number_format, c.font = fmt, SMALL
    chart = BarChart()
    chart.type = "bar"
    chart.title = (r.stub_label or r.stub)[:80]
    chart.y_axis.numFmt = "0%"
    chart.y_axis.scaling.min = 0
    data = Reference(ws, min_col=left + 1, max_col=left + n_cols, min_row=top + 2,
                     max_row=top + 2 + len(r.row_labels))
    cats = Reference(ws, min_col=left, min_row=top + 3, max_row=top + 2 + len(r.row_labels))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height, chart.width = 7.5, 15
    ws.add_chart(chart, f"{get_column_letter(left + n_cols + 2)}{top}")


def render_static(results: list, style: ReportStyle, path: str | Path, meta: dict) -> Path:
    wb = Workbook()
    entries = []
    taken: set[str] = {"Contents"}
    if style.layout == "one_sheet":
        ws = wb.create_sheet("Tables")
        row = 1
        for i, r in enumerate(results, start=1):
            entries.append((f"T{i}", r.stub_label or r.stub, "Tables"))
            row = _write_static_table(ws, row, f"T{i}", r, style)
        ws.column_dimensions["A"].width = 42
        ws.freeze_panes = "B1"
    else:
        for i, r in enumerate(results, start=1):
            name = _sheet_name(i, r.stub, taken)
            ws = wb.create_sheet(name)
            entries.append((f"T{i}", r.stub_label or r.stub, name))
            _write_static_table(ws, 1, f"T{i}", r, style)
            ws.column_dimensions["A"].width = 42
            ws.freeze_panes = "B4"
    _contents(wb, style, meta, entries)
    path = Path(path)
    wb.save(path)
    return path


# --- refreshable ------------------------------------------------------------------


@dataclass
class Criteria:
    """One COUNTIFS: (data column letter, value) pairs, all required."""
    pairs: list[tuple[str, int | str]]


def _countifs(c: Criteria) -> str:
    rng = lambda col: f"Data!${col}$2:${col}${DATA_ROWS + 1}"  # noqa: E731
    parts = []
    for col, value in c.pairs:
        parts += [rng(col), f'"{value}"' if isinstance(value, str) else str(value)]
    return f"COUNTIFS({','.join(parts)})"


def cell_terms(stub_col: str, codes: list[int] | None, banner: tuple[str, int] | None,
               filters: list[tuple[str, list[int]]]) -> list[Criteria]:
    """The COUNTIFS terms whose SUM is one cell's count. `codes` None = answered
    at all (the base). A filter or net over several codes is a sum of exact terms,
    never an array constant, so any combination stays correct."""
    stub_options = [[(stub_col, "<>")]] if codes is None else [[(stub_col, c)] for c in codes]
    banner_part = [[banner]] if banner else [[]]
    filter_options = [[(col, code) for code in fcodes] for col, fcodes in filters]
    terms = []
    for combo in product(stub_options, banner_part, *filter_options):
        pairs: list = []
        for piece in combo:
            pairs += piece if isinstance(piece, list) else [piece]
        terms.append(Criteria(pairs))
    return terms


def formula(terms: list[Criteria]) -> str:
    return "+".join(_countifs(t) for t in terms)


def render_refreshable(df, datamap, stubs: list[str], banner: list[str],
                       filter_: dict[str, list[int]] | None, nets: dict[str, list],
                       style: ReportStyle, path: str | Path, meta: dict) -> Path:
    """`df` holds every column the formulas read; `nets[stub]` is a list of
    (label, codes)."""
    wb = Workbook()
    used = list(dict.fromkeys([*stubs, *banner, *(filter_ or {})]))
    data = wb.create_sheet("Data")
    letter = {}
    for j, name in enumerate(used, start=1):
        letter[name] = get_column_letter(j)
        data.cell(row=1, column=j, value=name).font = BOLD
    for i, row in enumerate(df.select(used).iter_rows(), start=2):
        for j, value in enumerate(row, start=1):
            if value is not None:
                data.cell(row=i, column=j, value=value)
    data.freeze_panes = "A2"
    labels = wb.create_sheet("Labels")
    labels.append(["Variable", "Code", "Label"])
    for name in used:
        for code, lab in (datamap[name].values or {}).items():
            labels.append([name, code, lab])
    filters = [(letter[v], list(c)) for v, c in (filter_ or {}).items()]
    fmt = _pct_format(style.decimals)
    columns = [None] + [(b, code) for b in banner for code in datamap[b].valid_codes()]
    entries = []
    taken: set[str] = {"Contents", "Data", "Labels"}
    one = wb.create_sheet("Tables") if style.layout == "one_sheet" else None
    top = 1
    for i, stub in enumerate(stubs, start=1):
        v = datamap[stub]
        ws = one or wb.create_sheet(_sheet_name(i, stub, taken))
        if one is None:
            top = 1
        entries.append((f"T{i}", v.label or stub, ws.title))
        ws.cell(row=top, column=1, value=f"T{i} · {v.label or stub}").font = TITLE
        if filter_:
            ws.cell(row=top + 1, column=1, value="Filter: " + "; ".join(
                f"{datamap[f].label or f}: "
                + ", ".join(datamap[f].values.get(c, str(c)) for c in cs)
                for f, cs in filter_.items())).font = SMALL
        head, base = top + 2, top + 3
        ws.cell(row=head, column=1, value=stub).font = BOLD
        ws.cell(row=base, column=1, value="Base (n)").font = SMALL
        for j, col in enumerate(columns, start=2):
            label = "Total" if col is None else datamap[col[0]].values.get(col[1], str(col[1]))
            h = ws.cell(row=head, column=j, value=label)
            h.font, h.fill = BOLD, HEAD_FILL
            h.alignment = Alignment(wrap_text=True, horizontal="center")
            bc = (letter[col[0]], col[1]) if col else None
            ws.cell(row=base, column=j,
                    value="=" + formula(cell_terms(letter[stub], None, bc, filters))).font = SMALL
        row = base + 1
        rows = [(v.values.get(code, str(code)), [code], False) for code in v.valid_codes()]
        rows += [(label, list(codes), True) for label, codes in nets.get(stub, [])]
        for label, codes, is_net in rows:
            strong = is_net and style.bold_nets
            ws.cell(row=row, column=1, value=label).font = BOLD if strong else None
            for j, col in enumerate(columns, start=2):
                bc = (letter[col[0]], col[1]) if col else None
                base_ref = f"{get_column_letter(j)}${base}"
                c = ws.cell(row=row, column=j, value=(
                    f"=IFERROR(({formula(cell_terms(letter[stub], codes, bc, filters))})"
                    f"/{base_ref},0)"))
                c.number_format = fmt
                if is_net and style.bold_nets:
                    c.font, c.fill = BOLD, NET_FILL
            row += 1
        if style.charts:
            n_codes = len(v.valid_codes())
            _chart(ws, v.label or stub, base + 1, base + n_codes, len(columns),
                   f"{get_column_letter(len(columns) + 3)}{top}")
        ws.column_dimensions["A"].width = 42
        top = row + 2
    _contents(wb, style, {**meta, "notes": [
        "Refreshable: paste the next wave's respondents into the Data sheet, under the same "
        "column names (codes, not labels). Every base and percentage recalculates.",
        "Significance letters and means are in the finished version only.",
        *meta.get("notes", [])]}, entries)
    path = Path(path)
    wb.save(path)
    return path


if __name__ == "__main__":
    t = cell_terms("A", [1, 2], ("B", 3), [("C", [1, 2])])
    assert len(t) == 4 and t[0].pairs == [("A", 1), ("B", 3), ("C", 1)]
    assert formula(cell_terms("A", None, None, [])) == 'COUNTIFS(Data!$A$2:$A$100001,"<>")'
    print("ok")
