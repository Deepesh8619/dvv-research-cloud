"""The questionnaire, as a workbook somebody can read away from the screen.

Everything here is already known — preflight runs these checks before a link
can be made, and the review page shows each question. What it could not do is
leave the screen. A scripter working in another tool, a client signing off, a
PM checking a wave against the last one: all of them want the instrument as a
file, and the answer until now was a screenshot.

Built from the SurveyDefinition and the same audits that gate fielding, never
by re-reading the document. Two readings of one questionnaire that disagree is
worse than one reading, and the document has already been read.

Sheets, in the order somebody reads them:

    Summary        counts, the study brief, what the checks found
    Questions      one row each, with its base, type and flags
    Answers        every option, with its code and what is attached to it
    Routing        the conditions, as written and as compiled
    Checks         preflight, with the questions each finding names
    Flow           order, and what decides whether each question is seen
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.core.survey_definition import QuestionType, SurveyDefinition

_HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_TITLE_FONT = Font(size=12, bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")
_THIN = Border(bottom=Side(style="thin", color="DDDDDD"))
_MAX_WIDTH = 70
#: Cells Excel refuses to store, and a questionnaire can genuinely exceed it.
_CELL_LIMIT = 32000


def _sheet(workbook: Workbook, name: str, headers: list[str]):
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _WRAP
        cell.border = _THIN
    sheet.freeze_panes = "A2"
    return sheet


def _cell(value: Any) -> Any:
    """Anything openpyxl will accept, truncated where Excel would refuse it."""
    if value is None or isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    return text if len(text) <= _CELL_LIMIT else text[: _CELL_LIMIT - 1] + "…"


def _row(sheet, values: list[Any]) -> None:
    sheet.append([_cell(v) for v in values])


def _condition(expression) -> str:
    """What the script said, preferred over what we made of it.

    The verbatim is the thing a reviewer can check against their own document;
    the compiled form is only useful when there is no verbatim to show.
    """
    if expression is None:
        return ""
    return getattr(expression, "verbatim_text", "") or str(expression.to_dict())


def _base_of(question) -> str:
    return _condition(question.show_if) or "Everyone"


def build_workbook(
    survey: SurveyDefinition,
    path: str | Path,
    checks: list[dict[str, Any]] | None = None,
    brief: dict[str, Any] | None = None,
) -> Path:
    """Write the questionnaire workbook. Returns the path written."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    questions = list(survey.questions)
    asked = [q for q in questions if q.type is not QuestionType.INFO]
    checks = checks or []
    brief = brief or {}

    summary = _sheet(workbook, "Summary", ["", ""])
    summary["A1"].value = "Questionnaire"
    summary["B1"].value = survey.title or survey.study_id
    facts: list[tuple[str, Any]] = [
        ("Questions asked", len(asked)),
        ("Display screens", len(questions) - len(asked)),
        ("With a base condition", sum(1 for q in asked if q.show_if is not None)),
        ("With a screen-out", sum(1 for q in asked if q.terminate_if is not None)),
        ("Randomised", sum(1 for q in asked if q.randomize_options)),
    ]
    if brief.get("loi_minutes"):
        facts.append(("Stated length (minutes)", brief["loi_minutes"]))
    if brief.get("sample_size"):
        facts.append(("Stated sample size", brief["sample_size"]))
    if brief.get("markets"):
        facts.append(("Markets", brief["markets"]))
    for scope in brief.get("scope_statements") or []:
        facts.append(("What the study asks about", scope))
    for quota in brief.get("quotas") or []:
        facts.append((
            f"Quota — {quota.get('group', '')}",
            quota.get("definition") or quota.get("target_n") or "",
        ))
    failing = [c for c in checks if c.get("status") == "fail"]
    warning = [c for c in checks if c.get("status") == "warn"]
    facts.extend([("Checks failing", len(failing)), ("Checks to look at", len(warning))])
    for label, value in facts:
        _row(summary, [label, value])
    for cell in summary["A"]:
        cell.font = _TITLE_FONT if cell.row == 1 else Font(size=10, bold=True)

    detail = _sheet(workbook, "Questions", [
        "Order", "Question ID", "Type", "Question text", "Asked of",
        "Screen-out", "Answers", "Randomised", "Flags",
    ])
    for order, question in enumerate(questions, 1):
        _row(detail, [
            order, question.qid, str(question.type), question.text,
            _base_of(question), _condition(question.terminate_if),
            len(question.options or []) or len(question.rows or []),
            "Yes" if question.randomize_options else "",
            ", ".join(question.flags or ()),
        ])

    answers = _sheet(workbook, "Answers", [
        "Question ID", "Code", "Answer", "Row / item", "Attributes", "Shown if",
    ])
    for question in questions:
        for option in question.options or []:
            attributes = [
                name for name, on in (
                    ("EXCLUSIVE", option.exclusive), ("ANCHOR", option.anchored),
                ) if on
            ]
            _row(answers, [
                question.qid, option.code, option.label, "",
                ", ".join(attributes), _condition(option.show_if),
            ])
        for row in question.rows or []:
            _row(answers, [question.qid, row.code, row.label, "ROW", "", ""])

    routing = _sheet(workbook, "Routing", [
        "Question ID", "Rule", "As written in the document", "As programmed",
    ])
    for question in questions:
        for label, expression in (
            ("Asked of", question.show_if), ("Screen-out", question.terminate_if)
        ):
            if expression is not None:
                _row(routing, [
                    question.qid, label,
                    getattr(expression, "verbatim_text", "") or "",
                    str(expression.to_dict()),
                ])

    found = _sheet(workbook, "Checks", ["Status", "Check", "Detail", "Questions"])
    for check in checks:
        _row(found, [
            (check.get("status") or "").upper(), check.get("label") or check.get("id"),
            check.get("detail") or "", ", ".join(check.get("qids") or []),
        ])

    flow = _sheet(workbook, "Flow", [
        "Order", "Question ID", "Previous", "Next", "Seen when", "Leaves the survey when",
    ])
    for index, question in enumerate(questions):
        _row(flow, [
            index + 1, question.qid,
            questions[index - 1].qid if index else "",
            questions[index + 1].qid if index + 1 < len(questions) else "",
            _base_of(question), _condition(question.terminate_if),
        ])

    for sheet in workbook.worksheets:
        for column in range(1, sheet.max_column + 1):
            widest = max(
                (len(str(sheet.cell(r, column).value or ""))
                 for r in range(1, min(sheet.max_row, 200) + 1)),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(
                _MAX_WIDTH, max(10, widest + 2)
            )
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = _WRAP
        sheet.sheet_view.showGridLines = False

    path = Path(path)
    workbook.save(path)
    return path
