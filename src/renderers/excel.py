"""Excel tab deck renderer.

Renders a DeckRun to a formatted .xlsx: a summary sheet (TOC, weighting and
sig-test notes) plus one sheet per table. Layout follows tab-shop convention:
column letters under headers, weighted/unweighted base rows, category rows as
column %, bold NET rows, mean row, low-base columns flagged with *.
"""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.engines.tabs.crosstab import CrosstabResult
from src.engines.tabs.runner import DeckRun

HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
LABEL_FONT = Font(size=10)
BOLD_FONT = Font(size=10, bold=True)
NOTE_FONT = Font(size=9, italic=True, color="666666")
TITLE_FONT = Font(size=13, bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="DDDDDD"))
CENTER = Alignment(horizontal="center")


def _cell_text(pct: float, letters: str) -> str:
    return f"{pct:.1f}" + (f" {letters}" if letters else "")


def _write_table_sheet(ws: Worksheet, result: CrosstabResult, min_base: int) -> None:
    ws["A1"] = result.stub_label
    ws["A1"].font = TITLE_FONT
    subtitle = "Column %"
    if result.filter_desc:
        subtitle += f" — filtered to {result.filter_desc}"
    ws["A2"] = subtitle
    ws["A2"].font = NOTE_FONT

    header_row, letter_row, wbase_row, ubase_row = 4, 5, 6, 7
    first_data_row = 9

    ws.cell(row=header_row, column=1, value="").fill = HEADER_FILL
    for j, col in enumerate(result.columns):
        c = ws.cell(row=header_row, column=2 + j)
        flag = " *" if (col.letter and col.effective_base < min_base) else ""
        c.value = col.label + flag
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        letter_cell = ws.cell(row=letter_row, column=2 + j)
        letter_cell.value = f"({col.letter})" if col.letter else ""
        letter_cell.font = NOTE_FONT
        letter_cell.alignment = CENTER
        ws.cell(row=wbase_row, column=2 + j, value=round(col.base)).alignment = CENTER
        ws.cell(row=ubase_row, column=2 + j, value=col.unweighted_base).alignment = CENTER

    ws.cell(row=wbase_row, column=1, value="Base: weighted").font = NOTE_FONT
    ws.cell(row=ubase_row, column=1, value="Base: unweighted").font = NOTE_FONT

    row = first_data_row
    for i, label in enumerate(result.row_labels):
        ws.cell(row=row, column=1, value=label).font = LABEL_FONT
        for j in range(len(result.columns)):
            c = ws.cell(row=row, column=2 + j)
            c.value = _cell_text(result.col_pct[i, j], result.sig[i][j])
            c.alignment = CENTER
            c.border = THIN_BORDER
        row += 1

    if result.net_pct is not None:
        for gi, net_label in enumerate(result.net_labels):
            ws.cell(row=row, column=1, value=net_label).font = BOLD_FONT
            for j in range(len(result.columns)):
                c = ws.cell(row=row, column=2 + j)
                c.value = _cell_text(result.net_pct[gi, j], result.net_sig[gi][j])
                c.font = BOLD_FONT
                c.alignment = CENTER
            row += 1

    if result.means is not None:
        ws.cell(row=row, column=1, value="Mean").font = Font(size=10, italic=True, bold=True)
        for j in range(len(result.columns)):
            c = ws.cell(row=row, column=2 + j)
            letters = result.mean_sig[j] if result.mean_sig else ""
            c.value = f"{result.means[j]:.2f}" + (f" {letters}" if letters else "")
            c.alignment = CENTER
        row += 1

    ws.column_dimensions["A"].width = 34
    for j in range(len(result.columns)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 12
    ws.freeze_panes = ws.cell(row=first_data_row - 1, column=2)


def render_deck(run: DeckRun, path: str | Path) -> Path:
    """Write the deck to .xlsx and return the path."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    summary["A1"] = run.plan.title
    summary["A1"].font = Font(size=15, bold=True)
    summary["A2"] = f"Generated {datetime.now(UTC).strftime('%Y-%m-%d %H:%M UTC')}"
    summary["A2"].font = NOTE_FONT

    notes = [
        f"Weight: {run.plan.weight or 'unweighted'}",
        (
            f"Significance: column-proportions z-test at "
            f"{run.plan.confidence:.0%} confidence, computed on effective bases; "
            "a letter marks the column(s) the cell is significantly higher than."
        ),
        (
            f"Columns with effective base below {run.plan.min_base_for_test} are "
            "flagged * and excluded from testing."
        ),
    ]
    for i, note in enumerate(notes):
        cell = summary.cell(row=4 + i, column=1, value=note)
        cell.font = NOTE_FONT

    toc_start = 8
    summary.cell(row=toc_start, column=1, value="Table").font = BOLD_FONT
    summary.cell(row=toc_start, column=2, value="Title").font = BOLD_FONT
    summary.cell(row=toc_start, column=3, value="Filter").font = BOLD_FONT

    for idx, result in enumerate(run.results, start=1):
        sheet_name = f"T{idx}"
        ws = wb.create_sheet(sheet_name)
        _write_table_sheet(ws, result, run.plan.min_base_for_test)
        link_cell = summary.cell(row=toc_start + idx, column=1, value=sheet_name)
        link_cell.hyperlink = f"#'{sheet_name}'!A1"
        link_cell.font = Font(size=10, color="2B44CC", underline="single")
        summary.cell(row=toc_start + idx, column=2, value=result.stub_label).font = LABEL_FONT
        summary.cell(row=toc_start + idx, column=3, value=result.filter_desc or "—").font = (
            LABEL_FONT
        )

    summary.column_dimensions["A"].width = 12
    summary.column_dimensions["B"].width = 46
    summary.column_dimensions["C"].width = 36

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out
