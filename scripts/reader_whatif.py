"""What if the same questionnaire were written another way? Does the reader still read it?

A reader that only works on the documents it was built from is not ready for the
next client. This takes questionnaires whose reading is known (no client data:
point it at the synthetic set), rewrites each one in the styles real documents
use — "1." labels, "Q1)", "[Q1]", "Question 1:", the label on its own line, one
Word table per question, a translated line under the wording — and checks that
every question still comes back with the same id, type and answer codes as from
the original.

    python scripts/reader_whatif.py ~/dev/DVV/eval/synthetic/dev/*.docx
    python scripts/reader_whatif.py --style cards  some.docx     # one style only

Run with AI keys empty so only the deterministic reader is measured:
    BEDROCK_API_KEY= GROQ_API_KEY= OLLAMA_HOST=http://127.0.0.1:1 python scripts/...

A style below 100% is a reader gap with a name: the report lists the questions
that came back different and how. docs/READER_FIX_GUIDE.md says where to fix it.
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
import tempfile
from collections.abc import Callable
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402

from src.engines.questionnaire.extractor import extract_survey_definition  # noqa: E402
from src.intake.questionnaire_loader import load_docx, segment_into_questions  # noqa: E402

# "Q1.<tab>wording", "Q12a. wording" — how the synthetic set labels questions.
_LABEL = re.compile(r"^(Q(\d+)([a-z]?))[.)]\s+(.*)$", re.S)
# "a)<tab>Yes" — a lettered answer line.
_ANSWER = re.compile(r"^([a-z])\)\s+(.*)$", re.S)


def _relabel(fmt: Callable[[str, str, str], str]) -> Callable:
    """A style that rewrites only the label of each question line."""
    def rewrite(src: Document) -> Document:
        for p in src.paragraphs:
            if m := _LABEL.match(p.text):
                qid, num, sub, rest = m.groups()
                _set(p, fmt(qid, num + sub, rest))
        return src
    return rewrite


def _own_line(src: Document) -> Document:
    """"Q1." alone, the wording on the next line."""
    for p in list(src.paragraphs):
        if m := _LABEL.match(p.text):
            qid, _n, _s, rest = m.groups()
            _set(p, f"{qid}.")
            p.insert_paragraph_before(f"{qid}.")
            _set(p, rest)
    return src


def _translated(src: Document) -> Document:
    """A second-language line straight under each question's wording."""
    for p in list(src.paragraphs):
        if _LABEL.match(p.text):
            nxt = p._p.getnext()
            line = p._parent.add_paragraph("¿Versión en español de la pregunta?")
            if nxt is not None:
                nxt.addprevious(line._p)
    return src


def _cards(src: Document) -> Document:
    """One table per question: name / label / wording, codes | answers, notes."""
    paras = list(src.paragraphs)
    i = 0
    while i < len(paras):
        m = _LABEL.match(paras[i].text)
        if not m:
            i += 1
            continue
        qid, _n, _s, rest = m.groups()
        answers, notes, j = [], [], i + 1
        while j < len(paras) and paras[j].text.strip() and not _LABEL.match(paras[j].text):
            if a := _ANSWER.match(paras[j].text):
                answers.append((str(ord(a.group(1)) - 96), a.group(2)))
            else:
                notes.append(paras[j].text)
            j += 1
        table = src.add_table(rows=3 if answers else 2, cols=2)
        head = table.rows[0].cells[0].merge(table.rows[0].cells[1])
        head.text = f"QUESTION {qid}"
        head.add_paragraph(f"{qid}.")
        head.add_paragraph(rest)
        if answers:
            _lines(table.rows[1].cells[0], [c for c, _ in answers])
            _lines(table.rows[1].cells[1], [a for _, a in answers])
        foot = table.rows[-1].cells[0].merge(table.rows[-1].cells[1])
        _lines(foot, notes or ["Allow one response."])
        paras[i]._p.addprevious(table._tbl)
        for p in paras[i:j]:
            p._p.getparent().remove(p._p)
        i = j
    return src


def _lines(cell, lines: list[str]) -> None:
    cell.text = lines[0] if lines else ""
    for line in lines[1:]:
        cell.add_paragraph(line)


def _set(paragraph, text: str) -> None:
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


STYLES: dict[str, tuple[Callable, list[str]]] = {
    # name: (rewrite, labels an owner would teach for it — [] means none needed)
    "original": (lambda d: d, []),
    "paren": (_relabel(lambda q, n, r: f"{q})\t{r}"), []),
    "dash": (_relabel(lambda q, n, r: f"{q} – {r}"), []),
    "bracket": (_relabel(lambda q, n, r: f"[{q}] {r}"), ["[Q5]"]),
    "number_only": (_relabel(lambda q, n, r: f"{n}.\t{r}"), ["5."]),
    "question_word": (_relabel(lambda q, n, r: f"Question {n}: {r}"), ["Question 5:"]),
    "own_line": (_own_line, []),
    "translated": (_translated, []),
    "cards": (_cards, []),
}


def read(path: Path, taught: list[str]) -> dict[str, tuple]:
    from src.ai.router import BackendUnavailable

    blocks = load_docx(path)
    try:
        survey = extract_survey_definition(segment_into_questions(blocks, taught), "w",
                                           blocks=blocks)
    except BackendUnavailable:
        return {}  # the reader could not read it and fell to a model: nothing read
    return {q.qid.upper(): (q.type.value, tuple(o.code for o in q.options or []))
            for q in survey.questions if not q.qid.startswith(("INFO", "INTRO", "UNKNOWN"))}


def compare(base: dict, got: dict) -> tuple[int, list[str]]:
    misses = []
    for qid, (kind, codes) in base.items():
        if qid not in got:
            misses.append(f"{qid} not found")
        elif got[qid] != (kind, codes):
            misses.append(f"{qid} {kind}{list(codes)} -> {got[qid][0]}{list(got[qid][1])}")
    return len(base) - len(misses), misses


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("docs", nargs="+", type=Path)
    ap.add_argument("--style", action="append", choices=sorted(STYLES))
    ap.add_argument("--untaught", action="store_true",
                    help="do not apply the label an owner would teach (shows what needs teaching)")
    args = ap.parse_args(argv)
    logging.disable(logging.WARNING)
    styles = args.style or [s for s in STYLES if s != "original"]
    totals: dict[str, list[int]] = {s: [0, 0] for s in styles}
    with tempfile.TemporaryDirectory() as tmp:
        for doc in args.docs:
            base = read(doc, [])
            for style in styles:
                rewrite, teach = STYLES[style]
                out = Path(tmp) / f"{doc.stem}_{style}.docx"
                rewrite(Document(str(doc))).save(out)
                same, misses = compare(base, read(out, [] if args.untaught else teach))
                totals[style][0] += same
                totals[style][1] += len(base)
                if misses:
                    print(f"{doc.stem} [{style}] {same}/{len(base)}: " + "; ".join(misses[:6])
                          + (" …" if len(misses) > 6 else ""))
    print("\nstyle            same as original")
    worst = 100.0
    for style, (same, total) in totals.items():
        share = 100 * same / total if total else 100.0
        worst = min(worst, share)
        print(f"{style:<16} {same}/{total}  {share:.1f}%")
    return 0 if worst == 100.0 else 1


if __name__ == "__main__":
    sys.exit(main())
