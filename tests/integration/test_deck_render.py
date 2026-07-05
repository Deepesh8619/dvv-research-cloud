"""End-to-end: synthetic study -> tab plan -> Excel deck, read back and verified."""

import pytest
from openpyxl import load_workbook

from src.engines.tabs.runner import run_tab_plan
from src.engines.tabs.spec import load_tab_plan
from src.renderers.excel import render_deck
from src.utils.synthetic import build_synthetic

PLAN_YAML = """
title: Integration test deck
weight: weight
banner: [gender, age]
tables:
  - stub: q1
    nets: auto
    mean: true
  - stub: q2
    filter: {region: [1]}
"""


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("deck")
    parquet, datamap = build_synthetic(20_000, tmp / "study.parquet", verbose=False)
    plan_path = tmp / "plan.yaml"
    plan_path.write_text(PLAN_YAML)
    plan = load_tab_plan(plan_path, datamap)
    run = run_tab_plan(parquet, plan, datamap)
    out = render_deck(run, tmp / "deck.xlsx")
    return run, out


def test_workbook_structure(deck):
    run, out = deck
    wb = load_workbook(str(out))
    assert wb.sheetnames == ["Summary", "T1", "T2"]
    summary = wb["Summary"]
    assert summary["A1"].value == "Integration test deck"
    assert summary["B9"].value == "Q1. Agreement statement 1"


def test_table_sheet_content_matches_engine(deck):
    run, out = deck
    wb = load_workbook(str(out))
    t1 = wb["T1"]
    result = run.results[0]
    assert t1["A1"].value == result.stub_label
    assert t1["B4"].value == "Total"
    # first category row, Total column must equal the engine's percentage
    rendered = t1["B9"].value.split()[0]
    assert float(rendered) == pytest.approx(result.col_pct[0, 0], abs=0.05)
    # base rows present and plausible
    assert t1["B6"].value == round(result.columns[0].base)
    assert t1["B7"].value == result.columns[0].unweighted_base


def test_nets_and_mean_rendered(deck):
    run, out = deck
    wb = load_workbook(str(out))
    t1 = wb["T1"]
    labels = [t1.cell(row=r, column=1).value for r in range(9, 18)]
    assert any(v and v.startswith("Top 2 Box") for v in labels)
    assert any(v and v.startswith("Bottom 2 Box") for v in labels)
    assert "Mean" in labels


def test_filter_described_on_sheet(deck):
    run, out = deck
    wb = load_workbook(str(out))
    t2 = wb["T2"]
    assert "North" in t2["A2"].value


def test_sig_letters_present_in_deck(deck):
    # synthetic data has deliberate gender skew on q1 -> letters must appear
    run, out = deck
    result = run.results[0]
    all_letters = "".join("".join(row) for row in result.sig)
    assert "A" in all_letters or "B" in all_letters
