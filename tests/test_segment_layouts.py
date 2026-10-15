"""Question boundaries come from document structure, not from one house style.

Each fixture is a tiny invented document in a layout the parser used to read as
ONE question."""

from docx import Document

from src.intake.questionnaire_loader import load_docx, segment_into_questions


def _qids(path):
    return [c.candidate_qid for c in segment_into_questions(load_docx(path))]


def _spec_doc(tmp_path, labels):
    doc = Document()
    for qid in ("ROLE", "SIZE", "TOOL"):
        t = doc.add_table(rows=3, cols=2)
        values = (qid, "Single", f"Question {qid}?")
        for row, lab, val in zip(t.rows, labels, values, strict=True):
            row.cells[0].text, row.cells[1].text = lab, val
    p = tmp_path / "d.docx"
    doc.save(p)
    return p


def test_spec_tables_with_colon_labels(tmp_path):
    p = _spec_doc(tmp_path, ["Q Name:", "Q type:", "Question text:"])
    assert _qids(p) == ["ROLE", "SIZE", "TOOL"]


def test_spec_tables_with_variable_header(tmp_path):
    p = _spec_doc(tmp_path, ["Variable", "Question type", "Wording"])
    assert _qids(p) == ["ROLE", "SIZE", "TOOL"]


def test_quota_table_is_not_a_question(tmp_path):
    doc = Document()
    t = doc.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ("Cell", "Target (n)", "Definition"), strict=True):
        c.text = v
    p = tmp_path / "q.docx"
    doc.save(p)
    assert _qids(p) == [None]


def test_identifier_lines_without_q_numbers(tmp_path):
    doc = Document()
    for name in ("AGE", "MAIN_BANK", "LIFESTAGE (derived)"):
        doc.add_paragraph(name)
        doc.add_paragraph("(Single response) What is your answer to this?")
        doc.add_paragraph("[1] Yes")
    p = tmp_path / "i.docx"
    doc.save(p)
    assert _qids(p) == ["AGE", "MAIN_BANK", "LIFESTAGE"]


def test_lone_caps_word_is_not_a_question(tmp_path):
    doc = Document()
    for t in ("SCREENER", "Q1. First?", "Q2. Second?"):
        doc.add_paragraph(t)
    p = tmp_path / "l.docx"
    doc.save(p)
    assert _qids(p) == [None, "Q1", "Q2"]


def test_sub_item_numbering_with_spaces(tmp_path):
    doc = Document()
    for t in (
        "Q7_1  NUM 1-60 How many nights?", "Filter: Q3 in (1,3)",
        "Q7_2  NUM How many people?", "Q8  SA How booked?",
    ):
        doc.add_paragraph(t)
    p = tmp_path / "s.docx"
    doc.save(p)
    assert _qids(p) == ["Q7_1", "Q7_2", "Q8"]
