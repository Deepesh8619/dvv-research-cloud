"""The questionnaire as a file somebody can read away from the screen."""

from openpyxl import load_workbook

from src.core.survey_definition import (
    Block,
    Expression,
    Option,
    Question,
    QuestionType,
    SurveyDefinition,
)
from src.renderers.questionnaire_workbook import build_workbook


def _survey() -> SurveyDefinition:
    screener = Question(qid="S1", text="What is your age?", type=QuestionType.NUMERIC)
    brands = Question(
        qid="Q2",
        text="Which brands do you know?",
        type=QuestionType.MULTI,
        options=[
            Option(code=1, label="Alpha"),
            Option(code=2, label="Beta"),
            Option(code=3, label="None of these", exclusive=True),
        ],
        show_if=Expression.from_dict(
            {"op": "eq", "qid": "S1", "value": 1, "verbatim_text": "ASK IF S1=1"}
        ),
        randomize_options=True,
    )
    survey = SurveyDefinition(
        study_id="t",
        title="Demo study",
        blocks=[Block(block_id="b", question_ids=["S1", "Q2"])],
        questions=[screener, brands],
    )
    return survey


def _built(tmp_path, **kwargs):
    path = build_workbook(_survey(), tmp_path / "wb.xlsx", **kwargs)
    return load_workbook(path)


def test_every_sheet_a_reviewer_expects_is_there(tmp_path):
    workbook = _built(tmp_path)
    assert workbook.sheetnames == [
        "Summary", "Questions", "Answers", "Routing", "Checks", "Flow",
    ]


def test_each_question_and_each_answer_gets_a_row(tmp_path):
    workbook = _built(tmp_path)
    questions = list(workbook["Questions"].iter_rows(min_row=2, values_only=True))
    assert [r[1] for r in questions] == ["S1", "Q2"]
    answers = list(workbook["Answers"].iter_rows(min_row=2, values_only=True))
    assert [r[2] for r in answers] == ["Alpha", "Beta", "None of these"]
    assert answers[2][4] == "EXCLUSIVE"


def test_routing_shows_what_the_document_said_not_only_what_we_made_of_it(tmp_path):
    """The verbatim is the thing a reviewer can check against their own copy;
    the compiled form only helps when there is no verbatim to show."""
    workbook = _built(tmp_path)
    routing = list(workbook["Routing"].iter_rows(min_row=2, values_only=True))
    assert routing[0][0] == "Q2"
    assert routing[0][2] == "ASK IF S1=1"
    assert "S1" in routing[0][3]
    # And it reads as the base on the question row, not as a raw expression.
    questions = list(workbook["Questions"].iter_rows(min_row=2, values_only=True))
    assert questions[0][4] == "Everyone"
    assert questions[1][4] == "ASK IF S1=1"


def test_the_checks_travel_with_the_questionnaire(tmp_path):
    """A workbook that showed the instrument but not what is wrong with it
    would be the more dangerous half of the picture."""
    workbook = _built(
        tmp_path,
        checks=[
            {"id": "answers", "label": "Every question has answers",
             "status": "fail", "detail": "Q9 has nothing to pick", "qids": ["Q9"]},
        ],
    )
    rows = list(workbook["Checks"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0] == "FAIL"
    assert rows[0][3] == "Q9"
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Checks failing"] == 1


def test_the_study_brief_is_carried_onto_the_summary(tmp_path):
    workbook = _built(
        tmp_path,
        brief={
            "loi_minutes": 15,
            "scope_statements": ["Answer only about Greenbrook"],
            "quotas": [{"group": "Micro", "definition": "2-9 Employees"}],
        },
    )
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Stated length (minutes)"] == 15
    assert summary["What the study asks about"] == "Answer only about Greenbrook"
    assert summary["Quota — Micro"] == "2-9 Employees"


def test_a_cell_too_long_for_excel_is_truncated_not_refused(tmp_path):
    """A question's source text can exceed Excel's own cell limit, and a
    workbook that fails to save is worse than one with an ellipsis in it."""
    survey = _survey()
    survey.questions[0].text = "x" * 40000
    path = build_workbook(survey, tmp_path / "big.xlsx")
    value = load_workbook(path)["Questions"].cell(row=2, column=4).value
    assert len(value) < 40000 and value.endswith("…")
