"""Content the document shows as taken out must not reach a respondent.

A tracking study arrives as last wave's questionnaire with this wave's edits
on top. The Northwell Vivo wave-2 document carries 151 runs of red
strikethrough: brands dropped from lists, wording cut, and a whole question
struck out — all of it read as live content, so a removed answer would have
been offered and a deleted question asked.
"""

from docx import Document

from src.intake.revisions import live_cell_text, live_text

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _para(document, runs):
    """A paragraph of (text, struck) runs, built through python-docx."""
    paragraph = document.add_paragraph()
    for text, struck in runs:
        run = paragraph.add_run(text)
        run.font.strike = struck
    return paragraph


def test_a_struck_out_answer_is_not_part_of_the_question():
    document = Document()
    paragraph = _para(document, [("CVS, Walgreens, ", False), ("Vivo Health Pharmacy", True)])
    text, removed = live_text(paragraph)
    assert text == "CVS, Walgreens,"
    assert removed == ["Vivo Health Pharmacy"]


def test_a_line_struck_out_entirely_has_nothing_left():
    """A whole question struck out — the caller drops the block on empty text."""
    document = Document()
    paragraph = _para(document, [("Q13b. Why did you choose Amazon?", True)])
    text, removed = live_text(paragraph)
    assert text == ""
    assert removed == ["Q13b. Why did you choose Amazon?"]


def test_strikethrough_turned_off_again_is_live_text():
    """<w:strike w:val="false"/> is a real thing in a twice-edited document."""
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Kept after all")
    run.font.strike = False
    text, removed = live_text(paragraph)
    assert text == "Kept after all" and removed == []


def test_a_word_tracked_deletion_is_removed_like_a_strikethrough():
    """Word keeps deleted text in w:delText, so it contributes no visible text
    on its own — but it must still be REPORTED, or a reviewer cannot tell a
    deletion from something the parser missed."""
    document = Document()
    paragraph = document.add_paragraph()
    kept = paragraph.add_run("Which pharmacies ")
    deletion = paragraph._p.makeelement(f"{W}del", {})
    run = paragraph._p.makeelement(f"{W}r", {})
    deleted_text = paragraph._p.makeelement(f"{W}delText", {})
    deleted_text.text = "or someone you care for?"
    run.append(deleted_text)
    deletion.append(run)
    paragraph._p.append(deletion)

    text, removed = live_text(paragraph)
    assert text == kept.text.strip()
    assert removed == ["or someone you care for?"]


def test_a_struck_row_of_a_brand_list_is_dropped_from_the_cell():
    document = Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = ""
    run = cell.paragraphs[0].add_run("Vivo Health Pharmacy")
    run.font.strike = True
    text, removed = live_cell_text(cell)
    assert text == ""
    assert removed == ["Vivo Health Pharmacy"]
