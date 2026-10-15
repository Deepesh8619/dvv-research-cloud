"""Questionnaire loader: docx/PDF ingestion and rule-based segmentation.

Fixtures are generated programmatically (no real client files): a synthetic
.docx is authored with python-docx, and PDF-specific reading-order logic is
tested against a real synthetic PDF built with PyMuPDF plus direct unit tests of
the column-sorting helper. The shared `segment_into_questions` is also tested on
a hand-built block list so the segmentation rules are covered independently of
any file I/O.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from src.intake.questionnaire_loader import (
    DocumentBlock,
    _cluster_columns,
    _serialize_qname_table,
    _sort_reading_order,
    load_docx,
    load_pdf,
    segment_into_questions,
)

# --- fixtures ---------------------------------------------------------------


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Client Satisfaction Survey")  # preamble, no label
    doc.add_paragraph("Q1. What is your age?")

    section = doc.add_paragraph()
    run = section.add_run("SECTION B — SCREENER")
    run.bold = True

    doc.add_paragraph("S3. TERMINATE IF NOT AWARE")

    doc.add_paragraph("Q5. How much do you agree with each statement?")
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = ""
    table.cell(0, 1).text = "Agree"
    table.cell(1, 0).text = "Product is good value"
    table.cell(1, 1).text = "5"
    table.cell(2, 0).text = "Product is reliable"
    table.cell(2, 1).text = "4"

    out = tmp_path / "questionnaire.docx"
    doc.save(str(out))
    return out


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Q1. What is your age?", fontsize=11)
    page.insert_text((72, 130), "Q2. What is your gender?", fontsize=11)
    out = tmp_path / "questionnaire.pdf"
    doc.save(str(out))
    doc.close()
    return out


# --- docx loading + segmentation --------------------------------------------


def test_docx_loads_paragraphs_and_table_without_flattening(sample_docx: Path) -> None:
    blocks = load_docx(sample_docx)

    table_blocks = [b for b in blocks if b.block_type == "table"]
    assert len(table_blocks) == 1
    tbl = table_blocks[0]
    # The grid must be preserved as rows, not flattened into a paragraph.
    assert tbl.rows is not None
    assert tbl.rows[1] == ["Product is good value", "5"]
    assert tbl.rows[2] == ["Product is reliable", "4"]
    # page_number is None for docx.
    assert all(b.page_number is None for b in blocks)


def test_docx_bold_section_detected(sample_docx: Path) -> None:
    blocks = load_docx(sample_docx)
    section = next(b for b in blocks if b.text.startswith("SECTION B"))
    assert section.is_bold is True


def test_docx_segmentation_produces_expected_chunks(sample_docx: Path) -> None:
    blocks = load_docx(sample_docx)
    chunks = segment_into_questions(blocks)

    qids = [c.candidate_qid for c in chunks]
    # Preamble (None), Q1, SECTION (None), S3, Q5.
    assert qids == [None, "Q1", None, "S3", "Q5"]

    # The preamble chunk carries the title and is not dropped.
    assert "Client Satisfaction Survey" in chunks[0].chunk_text

    # The table belongs to the Q5 chunk and is retained as a table block.
    q5 = chunks[-1]
    assert q5.candidate_qid == "Q5"
    assert any(b.block_type == "table" for b in q5.source_blocks)


def test_docx_missing_file_raises() -> None:
    with pytest.raises(FileNotFoundError):
        load_docx("/nonexistent/questionnaire.docx")


# --- PDF loading ------------------------------------------------------------


def test_pdf_loads_and_segments(sample_pdf: Path) -> None:
    blocks = load_pdf(sample_pdf)
    assert all(b.page_number == 1 for b in blocks)

    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q1", "Q2"]
    assert chunks[0].provenance.startswith("page 1")


def test_pdf_missing_file_raises() -> None:
    with pytest.raises(FileNotFoundError):
        load_pdf("/nonexistent/questionnaire.pdf")


# --- reading-order helpers (unit, no I/O) -----------------------------------


def test_cluster_columns_detects_two_columns() -> None:
    # Left edges at ~50 and ~350 on a 600-wide page: two clusters.
    cols = _cluster_columns([50, 52, 350, 351], page_width=600)
    assert len(cols) == 2


def test_cluster_columns_merges_incidental_indent() -> None:
    # 50 and 60 are within the 15% gap (90px) on a 600 page: one column.
    cols = _cluster_columns([50, 60, 55], page_width=600)
    assert len(cols) == 1


def test_sort_reading_order_two_columns_column_then_row() -> None:
    # Left column (x0=50) rows at y=100,200; right column (x0=350) at y=100,200.
    parsed = [
        {"text": "R-top", "x0": 350, "y0": 100, "size": 11, "is_bold": False},
        {"text": "L-top", "x0": 50, "y0": 100, "size": 11, "is_bold": False},
        {"text": "R-bot", "x0": 350, "y0": 200, "size": 11, "is_bold": False},
        {"text": "L-bot", "x0": 50, "y0": 200, "size": 11, "is_bold": False},
    ]
    ordered = [b["text"] for b in _sort_reading_order(parsed, page_width=600)]
    # Full left column first, then full right column.
    assert ordered == ["L-top", "L-bot", "R-top", "R-bot"]


def test_sort_reading_order_single_column_top_to_bottom() -> None:
    parsed = [
        {"text": "b", "x0": 50, "y0": 200, "size": 11, "is_bold": False},
        {"text": "a", "x0": 50, "y0": 100, "size": 11, "is_bold": False},
    ]
    ordered = [b["text"] for b in _sort_reading_order(parsed, page_width=600)]
    assert ordered == ["a", "b"]


# --- segmentation on hand-built blocks --------------------------------------


def _para(text: str, idx: int) -> DocumentBlock:
    return DocumentBlock(block_type="paragraph", text=text, paragraph_index=idx)


def test_segment_instruction_marker_is_a_boundary() -> None:
    blocks = [
        _para("Q1. Intro question", 0),
        _para("SKIP TO Q5 IF UNDER 18", 1),
        _para("Q5. Main question", 2),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q1", None, "Q5"]
    # The SKIP TO instruction is its own chunk, not swallowed into Q1.
    assert "SKIP TO Q5" in chunks[1].chunk_text


def test_segment_qid_normalizes_case_and_suffix() -> None:
    blocks = [_para("a1a. subquestion", 0), _para("q12: another", 1)]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["A1a", "Q12"]


def test_segment_plain_word_is_not_a_label() -> None:
    blocks = [_para("Some intro text about the study", 0), _para("Q1. Real question", 1)]
    chunks = segment_into_questions(blocks)
    assert chunks[0].candidate_qid is None
    assert chunks[1].candidate_qid == "Q1"


def test_segment_arbitrary_letter_prefixes_split() -> None:
    """Pilot regression: a real screener used P1-P3 (parent path) and T1-T5
    (teen path); the old Q|S|A|C|D whitelist glued them onto the preceding
    question. Any short letter prefix + number must start a new chunk."""
    blocks = [
        _para("S6. Which stores have you shopped at?", 0),
        _para("P1. Are you the parent or guardian of any teenagers?", 1),
        _para("T1. How old are you?", 2),
        _para("PN2. Two-letter prefix question", 3),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["S6", "P1", "T1", "PN2"]


def test_segment_table_stays_with_preceding_question() -> None:
    tbl = DocumentBlock(
        block_type="table",
        text="Row A | 5",
        paragraph_index=1,
        rows=[["Row A", "5"]],
    )
    blocks = [_para("Q7. Grid question", 0), tbl, _para("Q8. Next", 2)]
    chunks = segment_into_questions(blocks)
    q7 = chunks[0]
    assert q7.candidate_qid == "Q7"
    assert any(b.block_type == "table" for b in q7.source_blocks)


# --- table-per-question format (real Google/GBK programming QRE shape) -------


def _qtable(text: str, idx: int) -> DocumentBlock:
    """A table block whose flattened text is passed verbatim (as load_docx would
    produce it) — merged cells repeat a value across a row, joined by ' | '."""
    return DocumentBlock(block_type="table", text=text, paragraph_index=idx, rows=[[text]])


def test_segment_table_per_question_each_qname_table_is_a_chunk() -> None:
    blocks = [
        _para("Study Overview", 0),  # preamble, no qid
        _qtable("INTRO | INTRO\nQ Name | TQ_INTRO | TQ_INTRO\nWelcome.", 1),
        _qtable("AGE | AGE\nQ Name | RESP_AGE | RESP_AGE\nWhat is your age?", 2),
        _qtable("EMP | EMP\nQName | SS_EMP_STATUS\nEmployment status?", 3),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == [
        None,
        "TQ_INTRO",
        "RESP_AGE",
        "SS_EMP_STATUS",  # 'QName' (no space) also matches
    ]
    # The preamble paragraph is its own headless chunk, not dropped.
    assert "Study Overview" in chunks[0].chunk_text


def test_segment_qname_merged_cell_takes_first_token() -> None:
    # Merged cells repeat the qid across the row; only the first is captured.
    blocks = [_qtable("Q Name | RESP_AGE | RESP_AGE | RESP_AGE\ntext", 0)]
    chunks = segment_into_questions(blocks)
    assert chunks[0].candidate_qid == "RESP_AGE"


def test_segment_table_without_qname_attaches_to_preceding_question() -> None:
    blocks = [
        _qtable("Q Name | TQ_BRAND | TQ_BRAND\nBrand awareness?", 0),
        # A quota/outline grid with no 'Q Name' row keeps the old behavior.
        _qtable("Quota Groups | Definition | Target N\nMicro | 2-9 | 75", 1),
        _qtable("Q Name | RESP_AGE\nAge?", 2),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["TQ_BRAND", "RESP_AGE"]
    # The Q-Name-less grid rides with the preceding TQ_BRAND chunk.
    assert len(chunks[0].source_blocks) == 2
    assert "Quota Groups" in chunks[0].chunk_text


def test_segment_scripter_note_rides_with_following_question() -> None:
    blocks = [
        _qtable("Q Name | TQ_INTRO\nWelcome.", 0),
        _para("SCRIPTER NOTE: randomize the options below", 1),
        _qtable("Q Name | TQ_BRAND\nBrand awareness?", 2),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["TQ_INTRO", "TQ_BRAND"]
    # The note belongs to TQ_BRAND (following), not TQ_INTRO (preceding).
    assert "SCRIPTER NOTE" not in chunks[0].chunk_text
    assert "SCRIPTER NOTE" in chunks[1].chunk_text


def test_segment_table_format_builds_from_docx(tmp_path: Path) -> None:
    """End-to-end via python-docx: real tables flatten to 'cell | cell' rows and
    each Q-Name table becomes its own chunk with a token qid."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Programming QRE")  # preamble
    for header, qid, prompt in (
        ("AGE", "RESP_AGE", "What is your age?"),
        ("EMP", "SS_EMP_STATUS", "Employment status?"),
        ("COMP", "HID_COMP_SIZE", "Company size?"),
    ):
        t = doc.add_table(rows=3, cols=3)
        for j in range(3):
            t.cell(0, j).text = header
            t.cell(1, j).text = "Q Name" if j == 0 else qid
        t.cell(2, 0).text = prompt
    out = tmp_path / "programming_qre.docx"
    doc.save(str(out))

    chunks = segment_into_questions(load_docx(out))
    assert [c.candidate_qid for c in chunks] == [
        None,
        "RESP_AGE",
        "SS_EMP_STATUS",
        "HID_COMP_SIZE",
    ]


# --- Q-Name table serialization (merged-cell cleanup + boilerplate strip) ----


def test_serialize_qname_table_dedups_merged_cells_and_strips_boilerplate() -> None:
    """The merged-cell duplication pattern from the real GBK QRE collapses to
    clean `Label: value` lines; boilerplate-only cells become blank."""
    rows = [
        ["AGE", "AGE", "AGE", "AGE"],  # section header: all identical -> once
        ["Q Name", "RESP_AGE", "RESP_AGE", "RESP_AGE"],
        ["What is your age?"] * 4,  # question text: once, unlabelled
        ["", "Response list:", "Response list:", "Response list:"],
        ["99", "Prefer not to answer", "Prefer not to answer", "Prefer not to answer"],
        # merged col B blank + boilerplate in C/D -> asked of everyone (blank)
        ["Among who", "", "If not All, specify:", "If not All, specify:"],
        ["Q type", "", "If Custom, specify:", "If Custom, specify:"],
        ["Randomize", "", "If Yes, details:", "If Yes, details:"],
        ["Other details", "TERMINATE if 3-99 selected", "TERMINATE if 3-99 selected", ""],
    ]
    out = _serialize_qname_table(rows)
    lines = out.splitlines()

    assert lines[0] == "AGE"  # section header emitted once, not "AGE | AGE | AGE"
    assert "Q Name: RESP_AGE" in lines
    assert "What is your age?" in lines  # label == value -> single line
    assert "Response list:" in lines  # empty label cell -> value alone
    assert "99: Prefer not to answer" in lines  # code: label
    assert "Among who: (blank)" in lines  # boilerplate-only -> blank
    assert "Q type: (blank)" in lines
    assert "Randomize: (blank)" in lines
    assert "Other details: TERMINATE if 3-99 selected" in lines
    # No merged-cell duplication anywhere and no boilerplate leaked through.
    assert "RESP_AGE | RESP_AGE" not in out
    assert "If not All, specify:" not in out
    assert "If Custom, specify:" not in out


def test_serialize_qname_table_keeps_genuinely_distinct_cells() -> None:
    """A true multi-column row (distinct, non-duplicate cells) is preserved,
    not collapsed — only merge duplicates are dropped."""
    out = _serialize_qname_table([["Scale", "Low", "Mid", "High"]])
    assert out == "Scale: Low | Mid | High"


def test_docx_qname_table_serialized_clean_end_to_end(tmp_path: Path) -> None:
    """The real .docx path (not just hand-built rows): a merged-cell Q-Name table
    loads with clean, deduplicated, boilerplate-free text."""
    from docx import Document

    doc = Document()
    t = doc.add_table(rows=5, cols=4)
    # Simulate merged cells by writing the same value across the row.
    for j in range(4):
        t.cell(0, j).text = "EMPLOYMENT"
        t.cell(1, j).text = "Q Name" if j == 0 else "SS_EMP_STATUS"
        t.cell(2, j).text = "1" if j == 0 else "Employed full-time"
    # Among who: blank value + boilerplate template in the trailing columns.
    t.cell(3, 0).text = "Among who"
    for j in (2, 3):
        t.cell(3, j).text = "If not All, specify:"
    # Other details: real terminate content, duplicated across the merge.
    t.cell(4, 0).text = "Other details"
    for j in range(1, 4):
        t.cell(4, j).text = "TERMINATE if 3-99 selected"
    out_path = tmp_path / "emp.docx"
    doc.save(str(out_path))

    blocks = load_docx(out_path)
    table = next(b for b in blocks if b.block_type == "table")

    assert "Q Name: SS_EMP_STATUS" in table.text
    assert "1: Employed full-time" in table.text
    assert "Among who: (blank)" in table.text
    assert "Other details: TERMINATE if 3-99 selected" in table.text
    # The merged-cell noise and boilerplate are gone.
    assert "SS_EMP_STATUS | SS_EMP_STATUS" not in table.text
    assert "If not All, specify:" not in table.text
    # Rows are still preserved for any downstream grid detection.
    assert table.rows is not None


# --- PDF: one block can hold several questions ------------------------------


@pytest.fixture
def tight_pdf(tmp_path: Path) -> Path:
    """A PDF whose lines sit close enough that PyMuPDF groups them into ONE block.

    This is the ordinary case in a real questionnaire, not a corner: consecutive
    lines of the same paragraph group become a single PDF "block", so the next
    question's label lands mid-block behind the previous question's options.
    """
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    for i, line in enumerate(
        [
            "Q1. What is your age?",
            "1. Under 30",
            "2. 30 or older",
            "Q2A. What is your gender?",
        ]
    ):
        page.insert_text((72, 100 + 13 * i), line, fontsize=11)
    out = tmp_path / "tight.pdf"
    doc.save(str(out))
    doc.close()
    return out


def test_pdf_block_holding_several_lines_is_split_per_line(tight_pdf: Path) -> None:
    import fitz

    with fitz.open(tight_pdf) as doc:
        text_blocks = [b for b in doc[0].get_text("dict")["blocks"] if b.get("type", 0) == 0]
    # Guard the premise: PyMuPDF really did merge all four lines into one block.
    assert len(text_blocks) == 1

    blocks = load_pdf(tight_pdf)
    assert [b.text for b in blocks] == [
        "Q1. What is your age?",
        "1. Under 30",
        "2. 30 or older",
        "Q2A. What is your gender?",
    ]


def test_pdf_question_label_after_options_starts_a_new_chunk(tight_pdf: Path) -> None:
    """The second question must not be swallowed by the first.

    Joining a block's lines into one string hid every label that was not at
    character 0; on the Federal Reserve SHED 2014 questionnaire that collapsed
    a ~100-question survey into 4 chunks.
    """
    chunks = segment_into_questions(load_pdf(tight_pdf))
    assert [c.candidate_qid for c in chunks] == ["Q1", "Q2A"]
    assert "Under 30" in chunks[0].chunk_text
    assert "Under 30" not in chunks[1].chunk_text


def test_uppercase_sub_question_suffix_is_a_boundary() -> None:
    """D4A/D8A (federal instruments) label as surely as A1a (Decipher docs)."""
    blocks = [
        DocumentBlock("paragraph", "D4. Are you employed?", 0),
        DocumentBlock("paragraph", "1. Yes", 1),
        DocumentBlock("paragraph", "D4A. Do you work full-time?", 2),
        DocumentBlock("paragraph", "A1a. And how many hours?", 3),
    ]
    chunks = segment_into_questions(blocks)
    # Suffix case is preserved as written — it is the data map's variable name.
    assert [c.candidate_qid for c in chunks] == ["D4", "D4A", "A1a"]


def test_base_tag_opens_a_question_when_nothing_is_numbered() -> None:
    """A questionnaire that numbers nothing and states a base per question.

    The healthcare/pharma convention. It matched no boundary rule — not a
    question label, not SECTION, not an instruction — so a whole instrument
    arrived as ONE headless chunk and the reviewer saw a single question whose
    text was the entire document. Measured on a real uploaded study: 1 question.
    """
    blocks = [
        DocumentBlock("paragraph", "Greenbrook TMS — TRD Treatment Landscape", 0),
        DocumentBlock("paragraph", "The survey will take approximately 15 minutes.", 1),
        DocumentBlock("paragraph", "[BASE: ALL] How many medications had they failed?", 2),
        DocumentBlock("paragraph", "After 1 antidepressant treatment", 3),
        DocumentBlock("paragraph", "After 2 antidepressant treatments", 4),
        DocumentBlock("paragraph", "[BASE: ALL] How satisfied are you with TMS?", 5),
        DocumentBlock("paragraph", "Very satisfied", 6),
        DocumentBlock("paragraph", "[BASE: THOSE WHO USE SPRAVATO] What drives the choice?", 7),
        DocumentBlock("paragraph", "Efficacy", 8),
    ]
    chunks = segment_into_questions(blocks)
    # The preamble keeps its own headless chunk; each base tag opens a question
    # and is NUMBERED, because a qid is the address review and approval use and
    # "UNKNOWN_3" is not one a reviewer can find in the document.
    assert [c.candidate_qid for c in chunks] == [None, "Q1", "Q2", "Q3"]
    assert "Very satisfied" in chunks[2].chunk_text
    assert "Very satisfied" not in chunks[3].chunk_text


def test_base_tag_under_a_numbered_question_is_not_a_boundary() -> None:
    """The rule that keeps every document that works today working.

    In a NUMBERED questionnaire the same tag annotates the question above it.
    Promoting it to a boundary there would split one question into two and
    orphan its options — so the promotion is gated on the document naming
    nothing at all. All three real client documents are numbered and contain no
    base-tag-opened lines, so none of them can change behaviour.
    """
    blocks = [
        DocumentBlock("paragraph", "Q5. How satisfied are you?", 0),
        DocumentBlock("paragraph", "[BASE: THOSE WHO USE SPRAVATO]", 1),
        DocumentBlock("paragraph", "Very satisfied", 2),
        DocumentBlock("paragraph", "Not satisfied", 3),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q5"]
    assert "Not satisfied" in chunks[0].chunk_text


def test_word_list_formatting_is_captured(tmp_path) -> None:
    """Word keeps list numbering in the XML, so the number never appears in the
    paragraph text — an answer list can arrive as bare lines with no visible
    codes at all. It is the document's own statement that a line is an ITEM
    rather than a sentence.

    Measured on the three real client documents: 322 of 598, 323 of 495 and 29
    of 55 paragraphs carry it, and it separates the answer list from the
    scripter notes beside it. Captured here; nothing reads it as a rule yet,
    because measuring showed a blanket "not a list item means not an option"
    would have destroyed two real answer lists that were typed plainly.
    """
    from docx import Document

    document = Document()
    document.add_paragraph("Q1. Which do you use?")
    document.add_paragraph("Alpha", style="List Number")
    document.add_paragraph("Beta", style="List Number")
    document.add_paragraph("[PN: RANDOMIZE]")
    path = tmp_path / "lists.docx"
    document.save(path)

    blocks = load_docx(path)
    listed = {b.text: b.is_list_item for b in blocks}
    assert listed["Alpha"] is True
    assert listed["Beta"] is True
    assert listed["Q1. Which do you use?"] is False
    assert listed["[PN: RANDOMIZE]"] is False


def test_a_question_whose_base_tag_was_forgotten_still_starts_its_own_question():
    """Both TRD studies had one: a question sentence straight after the previous
    question's answer table, without the [BASE:] tag every other question has.
    It rode inside the question before it and was never asked."""
    from src.intake.questionnaire_loader import DocumentBlock

    rows = [["Recommend one treatment", ""], ["Present several options", ""]]
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] How do you approach the referral? Select one.", 0),
        DocumentBlock("table", "\n".join(" | ".join(r) for r in rows), 1, rows=rows),
        DocumentBlock("paragraph", "Thinking about the patients you refer, what percentage "
                      "go to each setting? [0-100, MUST SUM TO 100%]", 2),
        DocumentBlock("paragraph", "[BASE: ALL] How satisfied are you overall?", 3),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q1", "Q2", "Q3"]
    assert chunks[1].chunk_text.startswith("Thinking about the patients you refer")


def test_a_note_after_an_answer_table_stays_with_its_question():
    from src.intake.questionnaire_loader import DocumentBlock

    rows = [["Yes", ""], ["No", ""]]
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] Do you prescribe TMS?", 0),
        DocumentBlock("table", "\n".join(" | ".join(r) for r in rows), 1, rows=rows),
        DocumentBlock("paragraph", "[TERMINATE IF NO — does this apply to referrers too?]", 2),
    ]
    assert len(segment_into_questions(blocks)) == 1


def test_lettered_sub_questions_are_questions_not_the_previous_ones_answers():
    """TRD referrer: "For the following questions..." then A.-D., each asking
    for a percentage, were all read as answers of the question before."""
    from src.intake.questionnaire_loader import DocumentBlock

    rows = [["After 1 treatment", ""], ["After 2 treatments", ""]]
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] When do you first discuss treatment? "
                      "Select one.", 0),
        DocumentBlock("table", "\n".join(" | ".join(r) for r in rows), 1, rows=rows),
        DocumentBlock("paragraph", "For the following questions, please provide your "
                      "best estimates.", 2),
        DocumentBlock("paragraph", "A. What percentage do you discuss treatment options with?", 3),
        DocumentBlock("paragraph", "Fill in [0-100]_____________ [%]", 4),
        DocumentBlock("paragraph", "B. Of those, what portion are referred for treatment?", 5),
        DocumentBlock("paragraph", "Fill in [0-100]_____________ [%]", 6),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q1", "Q2", "Q3"]
    assert chunks[1].chunk_text.startswith("A. What percentage")


def test_a_base_tag_written_without_a_colon_still_opens_a_question() -> None:
    """ "[BASE SELECTS 1-4 IN PREVIOUS QUESTION]" — no colon. Both TRD studies
    wrote it so, and the perception question rode inside the one above it."""
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] How familiar are you with X?", 0),
        DocumentBlock("paragraph", "Very familiar", 1),
        DocumentBlock("paragraph", "Not familiar", 2),
        DocumentBlock("paragraph", "[BASE SELECTS 1-4 IN PREVIOUS QUESTION] Which describes X?", 3),
        DocumentBlock("paragraph", "Hopeful", 4),
        DocumentBlock("paragraph", "[BASED ON YOUR ANSWERS] note text", 5),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q1", "Q2"]
    assert "Hopeful" not in chunks[0].chunk_text and "Hopeful" in chunks[1].chunk_text


def test_a_caps_heading_after_an_answer_table_is_not_an_answer() -> None:
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] Which describes X?", 0),
        DocumentBlock("table", "", 1, rows=[["Hopeful", ""], ["Unproven", ""]]),
        DocumentBlock("paragraph", "PRODUCT X PROFILE", 2),
        DocumentBlock("paragraph", "Please review the information below.", 3),
        DocumentBlock("paragraph", "[BASE: ALL] How strong is the evidence?", 4),
        DocumentBlock("table", "", 5, rows=[["Strong", ""]]),
        DocumentBlock("paragraph", "NONE OF THE ABOVE", 6),
    ]
    chunks = segment_into_questions(blocks)
    assert "PRODUCT X PROFILE" not in chunks[0].chunk_text
    assert "NONE OF THE ABOVE" in chunks[-1].chunk_text  # an answer, kept with its question


def test_a_bare_section_heading_is_not_a_screen() -> None:
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] Which describes X?", 0),
        DocumentBlock("table", "", 1, rows=[["Hopeful", ""]]),
        DocumentBlock("paragraph", "ADOPTION INTENT", 2),
        DocumentBlock("paragraph", "[BASE: ALL] How likely are you to use X?", 3),
        DocumentBlock("table", "", 4, rows=[["Very", ""]]),
    ]
    chunks = segment_into_questions(blocks)
    assert [c.candidate_qid for c in chunks] == ["Q1", "Q2"]
    assert all("ADOPTION INTENT" not in c.chunk_text for c in chunks)


def test_a_closing_thank_you_after_the_last_table_is_not_an_answer() -> None:
    blocks = [
        DocumentBlock("paragraph", "[BASE: ALL] What words describe X? Select up to 4.", 0),
        DocumentBlock("table", "", 1, rows=[["Novel", ""], ["Other (specify)", ""]]),
        DocumentBlock("paragraph", "Thank you! We appreciate your input!", 2),
    ]
    chunks = segment_into_questions(blocks)
    assert "Thank you" not in chunks[0].chunk_text
    assert chunks[-1].chunk_text.startswith("Thank you")
