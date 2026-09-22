"""The reading view: what the parser throws away, kept, and lined up with it."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from src.intake.document_html import render_docx
from src.intake.questionnaire_loader import load_docx


@pytest.fixture
def doc_path(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Client Satisfaction Survey")
    doc.add_paragraph("Q1. What is your age?")

    # A brand struck out of a live list: the LINE survives, the word does not.
    mixed = doc.add_paragraph()
    mixed.add_run("Q2. Which have you used? ")
    dropped = mixed.add_run("Bubble Co")
    dropped.font.strike = True
    mixed.add_run("Northwell")

    # A question struck out ENTIRELY: the loader drops it and does not count it.
    gone = doc.add_paragraph()
    run = gone.add_run("Q13b. Which pharmacy did you use last?")
    run.font.strike = True

    doc.add_paragraph("Q14. How satisfied were you?")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Brand"
    table.cell(0, 1).text = "Aware"
    table.cell(1, 0).text = "CVS"
    table.cell(1, 1).text = "Yes"

    out = tmp_path / "q.docx"
    doc.save(str(out))
    return out


def _ids(markup: str) -> list[str]:
    return re.findall(r'id="(b\d+)"', markup)


def test_every_anchor_matches_the_block_the_parser_extracted(doc_path: Path):
    """The load-bearing property: "para 47" must scroll to the right line.

    A question's provenance is the loader's block index. If this view counts
    blocks differently by so much as one, every anchor after that point sends
    the reviewer to the wrong paragraph — which is worse than no viewer,
    because it looks like the parser misread the document.
    """
    blocks = load_docx(doc_path)
    rendered = render_docx(doc_path)
    assert rendered.n_blocks == len(blocks)
    assert _ids(rendered.html) == [f"b{b.paragraph_index}" for b in blocks]


def test_a_wholly_struck_question_is_visible_and_unanchored(doc_path: Path):
    rendered = render_docx(doc_path)
    assert "Q13b. Which pharmacy did you use last?" in rendered.html
    # Rendered inside a block the loader never emitted, so it carries no id.
    gone_block = re.search(r'<p class="[^"]*dvv-gone[^"]*">(.*?)</p>', rendered.html)
    assert gone_block and "Q13b" in gone_block.group(1)


def test_a_brand_struck_out_of_a_live_list_is_marked_not_dropped(doc_path: Path):
    """The parser removes it from the answers; the reviewer still sees why."""
    rendered = render_docx(doc_path)
    assert "dvv-struck" in rendered.html
    assert "Bubble Co" in rendered.html
    assert "Northwell" in rendered.html


def test_a_table_is_a_table(doc_path: Path):
    rendered = render_docx(doc_path)
    assert "<table" in rendered.html and "<td>CVS</td>" in rendered.html


# --- untrusted input --------------------------------------------------------

def test_document_text_cannot_inject_markup(tmp_path: Path):
    """The document came from a client. Its text is data, never markup."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Q1. <script>alert('x')</script> & \"quoted\"")
    path = tmp_path / "x.docx"
    doc.save(str(path))

    rendered = render_docx(path)
    assert "<script>" not in rendered.html
    assert "&lt;script&gt;" in rendered.html
    assert "&amp;" in rendered.html


def test_only_a_real_hex_colour_reaches_the_style_attribute(tmp_path: Path):
    """A colour is emitted only when it is six hex digits — never echoed."""
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    paragraph = doc.add_paragraph()
    run = paragraph.add_run("coloured")
    colour = run._r.get_or_add_rPr().makeelement(qn("w:color"), {})
    colour.set(qn("w:val"), 'red" onload="alert(1)')
    run._r.get_or_add_rPr().append(colour)
    path = tmp_path / "c.docx"
    doc.save(str(path))

    rendered = render_docx(path)
    assert "onload" not in rendered.html
    assert "style=" not in rendered.html
