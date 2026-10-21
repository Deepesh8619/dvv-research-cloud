"""Findings read off the tables, each number cited; model sentences checked against cells."""

import polars as pl
from pptx import Presentation

from src.core.models import DataMap, Variable, VarType
from src.engines.analysis.findings import findings, unverified_numbers
from src.engines.tabs.runner import run_tab_plan
from src.engines.tabs.spec import TableSpec, TabPlan
from src.renderers.pptx import render_deck


def _table(tid, stub, title, rows, columns=None):
    columns = columns or [{"label": "Total", "letter": None, "base": 200.0, "unweighted_base": 200}]
    return {"id": tid, "stub": stub, "title": title, "columns": columns,
            "rows": [{"label": lbl, "pct": pct, "sig": sig or [""] * len(pct), "kind": "category"}
                     for lbl, pct, sig in rows], "means": None}


def test_leads_options_differences_and_the_sample_ranked_and_cited():
    tables = [
        _table("t1", "Country", "Country from URL", [("USA", [100.0], None), ("UK", [0.0], None)]),
        _table("t2", "S1", "Work status",
               [("Full-time", [93.0], None), ("Self-employed", [7.0], None)]),
        _table("t3", "Q5r1", "Website", [("Unchecked", [12.0], None), ("Checked", [88.0], None)]),
        _table("t4", "Q5r2", "Marketplace",
               [("Unchecked", [29.0], None), ("Checked", [71.0], None)]),
        _table("t5", "Q7", "Brand", [("Alpha", [40.0, 52.0, 25.0], ["", "B", ""]),
                                     ("Beta", [60.0, 48.0, 75.0], ["", "", "A"])],
               columns=[{"label": "Total", "letter": None, "base": 200.0, "unweighted_base": 200},
                        {"label": "Under 35", "letter": "A", "base": 100.0, "unweighted_base": 100},
                        {"label": "35+", "letter": "B", "base": 100.0, "unweighted_base": 24}]),
    ]
    found = findings(tables, {"Q5": "Where do you sell?"}, screeners={"S1"})
    kinds = [f["kind"] for f in found]
    # The tested difference first, the screener last, the constant country nowhere.
    assert kinds == ["difference", "most_chosen", "sample"]
    diff, multi, sample = found
    assert "52% among Under 35 against 25% among 35+" in diff["text"]
    assert diff["caution"] == "Base under 30 for 35+: indicative only."
    assert multi["text"] == "Where do you sell?: most chosen Website 88%, then Marketplace 71%."
    assert [(c["table"], c["value"], c["base"]) for c in multi["cites"]] == [
        ("t3", 88.0, 200), ("t4", 71.0, 200)]
    assert sample["qid"] == "S1"
    # Every computed finding verifies against its own tables.
    by_id = {t["id"]: t for t in tables}
    assert all(not unverified_numbers(f["text"], [by_id[i] for i in f["tables"]]) for f in found)


def test_a_number_not_in_the_table_is_caught():
    t = _table("t1", "Q1", "Q1", [("Yes", [61.4], None), ("No", [38.6], None)])
    assert unverified_numbers("61% said yes (n=200), 23 points ahead", [t]) == []
    assert unverified_numbers("73% said yes, up 12 points from n=180", [t]) == [
        "73%", "12points", "180"]
    coded = _table("t2", "Q2", "Q2", [("1", [71.0], None), ("2", [29.0], None)])
    assert unverified_numbers("71% chose 1", [coded]) == []
    assert unverified_numbers("72% chose 1", [coded]) == ["72%"]


def test_the_deck_opens_with_cited_findings_and_needs_no_model(tmp_path):
    pl.DataFrame({"Q1": [1] * 70 + [2] * 30}).write_parquet(tmp_path / "d.parquet")
    dm = DataMap([Variable(name="Q1", label="Heard of us?", type=VarType.SINGLE,
                           values={1: "Yes", 2: "No"}, question_ref="Q1")])
    run = run_tab_plan(tmp_path / "d.parquet", TabPlan(title="T", banner=[], tables=[
        TableSpec(stub="Q1")]), dm)
    slides = Presentation(render_deck(run, tmp_path / "d.pptx", narrate=False)).slides
    text = ["\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
            for slide in slides]
    assert "Key findings" in text[1]
    assert "Yes leads at 70%, ahead of No at 30%.  [T1: Total n=100]" in text[1]
    assert "Source: T1, base n=100" in text[2]
