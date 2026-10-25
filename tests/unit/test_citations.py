"""The report checker: numbers, comparisons, significance and small bases."""

import polars as pl
from pptx import Presentation

from src.core.models import DataMap, Variable, VarType
from src.engines.analysis.citations import check, to_cites
from src.engines.tabs.runner import run_tab_plan
from src.engines.tabs.spec import TableSpec, TabPlan
from src.renderers.pptx import render_deck

COLS = [{"label": "Total", "letter": None, "base": 200.0, "unweighted_base": 200},
        {"label": "18-34", "letter": "A", "base": 100.0, "unweighted_base": 100},
        {"label": "55+", "letter": "B", "base": 20.0, "unweighted_base": 20}]
T = {"id": "t12", "stub": "Q7", "title": "Brand", "columns": COLS, "means": None, "rows": [
    {"label": "Brand A", "pct": [30.0, 34.0, 22.0], "sig": ["", "B", ""], "kind": "category"},
    {"label": "Brand B", "pct": [70.0, 66.0, 78.0], "sig": ["", "", "A"], "kind": "category"}]}
TABLES = [T]
A, B = {"table": "t12", "row": 0, "col": 1}, {"table": "t12", "row": 0, "col": 2}


def test_a_correct_sentence_passes():
    assert check("Brand A reaches 34% of 18-34s.", [A], TABLES) == {"ok": True, "problems": []}
    assert check("Brand A leads among 18-34s at 34%, ahead of 22%, a 12 point gap, "
                 "indicative as 55+ is small.", [A, B], TABLES)["ok"]


def test_a_wrong_number_fails_with_a_plain_sentence():
    r = check("Brand A reaches 41% of 18-34s.", [A], TABLES)
    assert not r["ok"] and r["problems"] == [
        "41% matches none of the cited cells. Cite the cell it comes from or change the number."]


def test_a_comparison_needs_two_cells():
    r = check("Brand A is higher than average at 34%.", [A], TABLES)
    assert not r["ok"] and "fewer than two cells" in r["problems"][0]


def test_significantly_needs_the_sig_letter():
    ok = check("Brand A is significantly higher among 18-34s (34%) than 55+ (22%), "
               "indicative for 55+.", [A, B], TABLES)
    assert ok["ok"], ok
    other = {"table": "t12", "row": 1, "col": 2}  # a different row: no test between these
    r = check("Brand A (34%) is significantly ahead of Brand B (78%).", [A, other], TABLES)
    assert not r["ok"] and any("significant" in p for p in r["problems"])


def test_a_small_base_must_be_called_indicative():
    r = check("Brand A reaches 22% of 55+.", [B], TABLES)
    assert not r["ok"] and "Base under 30 for 55+ (n=20): say it is indicative." in r["problems"]
    assert check("Brand A reaches 22% of 55+, which is indicative.", [B], TABLES)["ok"]


def test_a_missing_cell_is_a_problem_not_a_crash():
    r = check("Brand A is 34%.", [{"table": "t99", "row": 0, "col": 0}], TABLES)
    assert not r["ok"] and "does not exist" in r["problems"][0]


def test_finding_cites_become_indexes():
    f = [{"table": "t12", "row": "Brand A", "column": "55+", "value": 22.0, "base": 20}]
    assert to_cites(f, TABLES) == [B]


def test_the_deck_has_a_report_slide_with_citations_in_the_notes(tmp_path):
    pl.DataFrame({"Q1": [1] * 70 + [2] * 30}).write_parquet(tmp_path / "d.parquet")
    dm = DataMap([Variable(name="Q1", label="Heard?", type=VarType.SINGLE,
                           values={1: "Yes", 2: "No"}, question_ref="Q1")])
    run = run_tab_plan(tmp_path / "d.parquet", TabPlan(title="T", banner=[], tables=[
        TableSpec(stub="Q1")]), dm)
    report = {"sentences": [
        {"text": "Yes leads at 70%.", "ok": True, "cites": [{"table": "t1", "row": 0, "col": 0}]},
        {"text": "Refused one.", "ok": False, "cites": []}]}
    slides = Presentation(render_deck(run, tmp_path / "d.pptx", narrate=False,
                                      report=report)).slides
    slide = slides[1]
    body = "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Report" in body and "Yes leads at 70%." in body and "Refused" not in body
    assert "T1 row 1, column 1" in slide.notes_slide.notes_text_frame.text


def test_a_multi_select_finding_cites_the_checked_row_of_each_option_table():
    from src.engines.analysis.citations import check, to_cites

    def option(tid, pct):
        return {"id": tid, "stub": f"Q5r{tid[-1]}", "title": "Q5", "filter": "",
                "columns": [{"label": "Total", "letter": "A", "base": 34, "unweighted_base": 34,
                             "effective_base": 34}],
                "rows": [{"label": "Checked", "pct": [pct], "sig": [""], "kind": "category"},
                         {"label": "Not checked", "pct": [100 - pct], "sig": [""],
                          "kind": "category"}],
                "means": None, "mean_sig": None}
    tables = [option("t1", 85.3), option("t2", 67.6)]
    cites = to_cites([{"table": "t1", "row": "Nescafé", "column": "Total"},
                      {"table": "t2", "row": "Bru", "column": "Total"}], tables)
    assert cites == [{"table": "t1", "row": 0, "col": 0}, {"table": "t2", "row": 0, "col": 0}]
    assert check("Most chosen Nescafé 85%, then Bru 68%.", cites, tables)["ok"]
